import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main import main
from src.excel_reader import load_config
from src.file_manager import validate_output_path


class StageTwoTests(unittest.TestCase):
    def make_queue(self, root, dry_run="SIM", limit=2):
        input_dir = root / "input"
        input_dir.mkdir()
        references = input_dir / "referencias"
        references.mkdir()
        (references / "ok.png").touch()
        output = root / "output" / "resultados"
        output.mkdir(parents=True)
        queue = input_dir / "fila.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Fila_Geracao"
        sheet.append(["ID", "Prompt_Padrao", "Prompt_Variacao", "Arquivo_Referencia",
                      "Nome_Saida", "Status", "Tentativas", "Observacao"])
        for i in range(1, 5):
            sheet.append([i, "base", "var", "ok.png", f"{i}.png", "PENDENTE", 0, None])
        sheet.append([5, "base", "var", "ok.png", "5.png", "CONCLUIDO", 1, None])
        config = workbook.create_sheet("Configuracao")
        for pair in [("Parametro", "Valor"), ("Pasta_Referencias", "referencias"),
                     ("Pasta_Resultados", "resultados"), ("Limite_Por_Execucao", limit),
                     ("Dry_Run", dry_run)]:
            config.append(pair)
        workbook.save(queue)
        workbook.close()
        return queue, output

    def test_config_and_dry_run_limit(self):
        with tempfile.TemporaryDirectory() as folder:
            queue, output = self.make_queue(Path(folder))
            config = load_config(queue)
            self.assertEqual(config["result_dir"], output.resolve())
            self.assertEqual(config["reference_dir"], (queue.parent / "referencias").resolve())
            stream = io.StringIO()
            with contextlib.redirect_stdout(stream):
                summary = main(queue)
            self.assertEqual(summary, {"processed": 2, "success": 2, "errors": 0, "ignored": 3})
            self.assertIn("DRY RUN | ID: 1", stream.getvalue())
            self.assertIn("Prompt final: base\n\nvar", stream.getvalue())
            self.assertIn("Nome do arquivo de saída: 2.png", stream.getvalue())
            self.assertNotIn("DRY RUN | ID: 3", stream.getvalue())
            book = load_workbook(queue)
            self.assertEqual(book["Fila_Geracao"]["F2"].value, "PENDENTE")
            self.assertEqual(book["Fila_Geracao"]["G2"].value, 0)
            book.close()

    def test_output_validation_and_error_continuation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            queue, output = self.make_queue(root, dry_run="NAO", limit=3)
            (output / "1.png").touch()
            for name in ("", "../bad.png", "a/b.png", "CON.png", "bad?.png", "trailing."):
                with self.assertRaises(ValueError):
                    validate_output_path(output, name)
            with self.assertRaises(FileExistsError):
                validate_output_path(output, "1.png")
            summary = main(queue)
            self.assertEqual(summary, {"processed": 3, "success": 2, "errors": 1, "ignored": 2})
            book = load_workbook(queue)
            sheet = book["Fila_Geracao"]
            self.assertEqual([sheet[f"F{i}"].value for i in range(2, 6)],
                             ["ERRO", "CONCLUIDO", "CONCLUIDO", "PENDENTE"])
            self.assertEqual(book["Configuracao"]["B3"].value, "resultados")
            book.close()
            self.assertEqual(list(output.iterdir()), [output / "1.png"])

    def test_invalid_limit(self):
        with tempfile.TemporaryDirectory() as folder:
            queue, _ = self.make_queue(Path(folder), limit=0)
            with self.assertRaises(ValueError):
                load_config(queue)


if __name__ == "__main__":
    unittest.main()
