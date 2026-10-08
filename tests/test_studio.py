import base64
from copy import deepcopy
from dataclasses import asdict, replace
import hashlib
import http.client
import io
import json
import os
from pathlib import Path
import tempfile
import threading
import time
from types import SimpleNamespace
from unittest.mock import patch, Mock

from PIL import Image
from openpyxl import load_workbook
from main import main, make_job
from src.excel_reader import load_config, load_queue
from src.execution_state import Journal, verified_outputs, digest
from src.image_generator import OpenAIImageGenerator, GenerationError
from src.image_settings import ImageSettings
from src.studio_session import StudioSession, validate_folders, validate_photos
from studio import make_server
from support import IsolatedTest, png_bytes
from test_stage_three import StageThreeTests


def image_bytes(fmt):
    stream = io.BytesIO()
    Image.new("RGB", (16, 16), "blue").save(stream, format=fmt.upper())
    return stream.getvalue()


class StudioTests(IsolatedTest):
    def setUp(self):
        super().setUp()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        rows = [[i, "Fotografia de produto", f"Variação {i}", "ok.png", f"{i}.png", "PENDENTE", 0, ""] for i in range(1, 9)]
        self.queue, self.refs, self.output = StageThreeTests().make_queue(self.root, rows)
        self.session = StudioSession(self.root / "settings.json", env={})
        self.session.choose_queue(self.queue)
        self.env = dict(IMAGE_PROVIDER="openai", DRY_RUN="NAO", FIRST_RUN_SAFE_MODE="NAO", MAX_JOBS_PER_RUN="20", MAX_IMAGES_PER_RUN="20")
        self.api = Mock(side_effect=lambda **kw: SimpleNamespace(data=[SimpleNamespace(
            b64_json=base64.b64encode(image_bytes(kw['output_format'])).decode()) for _ in range(kw['n'])]))

    def generator(self, settings=None):
        return OpenAIImageGenerator(client=SimpleNamespace(images=SimpleNamespace(edit=self.api)), settings=settings)

    def test_api_key_session_and_protected_storage_never_enter_state(self):
        if os.name != 'nt':
            self.skipTest('DPAPI is Windows-only')
        key = 'sk-test-placeholder-only-1234567890'
        external = 'sk-other-placeholder-only-1234567890'
        self.assertEqual(self.session.api_key_status(), 'missing')
        self.session.set_api_key(key)
        self.assertEqual(self.session.api_key_status(), 'session')
        self.assertEqual(self.session.resolve_api_key(external), key)
        self.assertNotIn(key, json.dumps(self.session.state()))
        self.assertFalse(self.session.credential_store.exists())
        self.session.set_api_key(key, persist=True)
        self.assertEqual(self.session.api_key_status(), 'saved')
        self.assertNotIn(key.encode(), self.session.credential_store.path.read_bytes())
        reopened = StudioSession(self.session.settings_path, env={'OPENAI_API_KEY': external})
        self.assertEqual(reopened.api_key_status(), 'saved')
        self.assertEqual(reopened.resolve_api_key(external), key)
        self.assertNotIn(key, json.dumps(reopened.state()))
        reopened.clear_api_key()
        self.assertEqual(reopened.api_key_status(), 'external')
        self.assertFalse(reopened.credential_store.exists())

    def edit_queue(self, fn):
        book = load_workbook(self.queue)
        try:
            fn(book)
            book.save(self.queue)
        finally:
            book.close()

    def test_old_workbook_and_session_selection_do_not_write_workbook(self):
        before = digest(self.queue)
        self.assertEqual(self.session.validate_campaign()['settings']['output_format'], 'png')
        self.session.set_mode('direct')
        self.session.add_photos([self.refs / 'ok.png'])
        self.assertEqual(len(self.session.validate_campaign()['jobs']), 8)
        self.assertEqual(digest(self.queue), before)
        self.assertFalse(Path(str(self.queue) + '.state.json').exists())

    def test_spreadsheet_selection_reports_prompt_and_image_counts_read_only(self):
        rows = [
            [1, "Prompt", "", "ok.png", "1.png", "PENDENTE", 0, "", "", "", "", "", 2, ""],
            [2, "", "Variation", "ok.png", "2.png", "PENDENTE", 0, "", "", "", "", "", None, ""],
            [3, "Prompt", "", "ok.png", "3.png", "CONCLUIDO", 0, "", "", "", "", "", 3, ""],
            [4, "Prompt", "", "ok.png", "4.png", "REVISAO", 0, "", "", "", "", "", 1, ""],
            [5, "", "", "ok.png", "5.png", "PENDENTE", 0, "", "", "", "", "", 5, ""],
        ]
        queue, _, _ = StageThreeTests().make_queue(self.root / "summary", rows, optional=True)
        before = digest(queue)
        self.session.choose_queue(queue)
        self.assertEqual(self.session.state()["queue_summary"], {
            "items_with_prompt": 4,
            "pending_items": 2,
            "estimated_images": 3,
            "execution_items": 1,
            "execution_images": 1,
            "item_limit": 1,
            "image_limit": 1,
            "invalid_quantities": 0,
            "completed_items": 1,
            "completed_images": 3,
            "review_items": 1,
        })
        self.assertEqual(digest(queue), before)

        book = load_workbook(queue)
        try:
            book["Fila_Geracao"]["M2"] = "inválida"
            book.save(queue)
        finally:
            book.close()
        self.session.choose_queue(queue)
        summary = self.session.state()["queue_summary"]
        self.assertEqual(summary["invalid_quantities"], 1)
        self.assertEqual(summary["estimated_images"], 1)
        preferences = deepcopy(self.session.preferences)
        preferences.update(max_jobs=5, max_images=2)
        self.session.set_preferences(preferences)
        updated = self.session.state()["queue_summary"]
        self.assertEqual((updated["execution_items"], updated["execution_images"]), (1, 1))
    def test_direct_ignores_invalid_sheet_references_and_missing_column(self):
        self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('D2', 'UNKNOWN::../../missing.png'))
        self.session.set_mode('direct')
        self.session.add_photos([self.refs / 'ok.png'])
        plan = self.session.validate_campaign()
        self.assertTrue(all(job['references'] == [str(self.refs / 'ok.png')] for job in plan['jobs']))
        self.edit_queue(lambda b: b['Fila_Geracao'].delete_cols(4))
        self.assertEqual(len(self.session.validate_campaign()['jobs']), 8)
        result = main(self.queue, env={**self.env, 'DRY_RUN': 'SIM'}, overrides=self.session.snapshot()[0])
        self.assertEqual(result['success'], 8)
        with self.assertRaisesRegex(ValueError, 'Colunas obrigatórias'):
            load_queue(self.queue)

    def test_sheet_mode_never_uses_session_photos(self):
        self.session.add_photos([self.refs / 'ok.png'])
        self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('D2', ''))
        with self.assertRaisesRegex(ValueError, 'Arquivo_Referencia'):
            self.session.validate_campaign()
        self.assertEqual(self.api.call_count, 0)

    def test_selection_cumulative_same_basename_remove_limit_and_content(self):
        other = self.root / 'second'
        other.mkdir()
        (other / 'ok.png').write_bytes(png_bytes())
        self.session.add_photos([self.refs / 'ok.png'])
        self.session.add_photos([other / 'ok.png'])
        self.session.add_photos([other / 'ok.png'])
        self.assertEqual(len(self.session.photos), 2)
        self.session.remove_photo(0)
        self.assertEqual(self.session.photos, [(other / 'ok.png').resolve()])
        for i in range(16):
            p = other / f'{i}.png'
            p.write_bytes(png_bytes())
        with self.assertRaisesRegex(ValueError, '16'):
            self.session.add_photos(list(other.glob('*.png')))
        self.assertEqual(len(self.session.photos), 1)
        invalid = other / 'fake.png'
        invalid.write_text('not an image')
        with self.assertRaisesRegex(ValueError, 'Conteúdo'):
            validate_photos([invalid])
        (other / 'ok.png').unlink()
        self.session.set_mode('direct')
        with self.assertRaises(FileNotFoundError):
            self.session.validate_campaign()

    def test_two_folders_same_name_and_explicit_atomic_save(self):
        second = self.root / 'brand'
        second.mkdir()
        (second / 'ok.png').write_bytes(image_bytes('png'))
        folders = [{'alias': 'Produtos', 'path': str(self.refs)}, {'alias': 'Marca', 'path': str(second)}]
        before = digest(self.queue)
        self.session.set_folders(folders, 'Produtos')
        self.assertEqual(digest(self.queue), before)
        self.session.save_folders()
        config = load_config(self.queue)
        self.assertEqual(config['reference_dir'], self.refs.resolve())
        self.assertEqual(config['reference_dirs']['Marca'], second.resolve())
        self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('D2', 'Produtos::ok.png; Marca::ok.png'))
        job = make_job(load_queue(self.queue).iloc[0], self.session.snapshot()[0], True)
        self.assertEqual(job.referencias, (self.refs / 'ok.png', second / 'ok.png'))
        self.assertEqual(self.session.folder_usage('Marca'), ['1'])
        with self.assertRaisesRegex(ValueError, 'padrão'):
            self.session.set_folders([folders[1]], '')
        with self.assertRaisesRegex(ValueError, 'duplicado'):
            validate_folders(folders + [{'alias': 'marca', 'path': str(second)}], 'Produtos')
        with self.assertRaisesRegex(ValueError, 'inválido'):
            validate_folders([{'alias': '../bad', 'path': str(second)}], '')

    def test_alias_unknown_traversal_and_link_escape(self):
        from src.file_manager import validate_reference_paths
        with self.assertRaisesRegex(ValueError, 'cadastrada'):
            validate_reference_paths(self.refs, 'Unknown::ok.png', {})
        with self.assertRaisesRegex(ValueError, 'fora'):
            validate_reference_paths(self.refs, '../ok.png')
        outside = self.root / 'outside'
        outside.mkdir()
        (outside / 'ok.png').write_bytes(png_bytes())
        link = self.refs / 'escape'
        try:
            link.symlink_to(outside, target_is_directory=True)
        except OSError:
            import os, subprocess
            if os.name != 'nt':
                self.skipTest('Symlink unsupported on this filesystem')
            subprocess.run(['cmd', '/c', 'mklink', '/J', str(link), str(outside)], check=True, capture_output=True)
        try:
            with self.assertRaisesRegex(ValueError, 'fora'):
                validate_reference_paths(self.refs, 'escape/ok.png')
        finally:
            if link.is_symlink():
                link.unlink()
            else:
                link.rmdir()  # remove junction itself, never recurse into the target.

    def test_settings_constraints_and_legacy_defaults(self):
        self.assertEqual(ImageSettings.from_env({}), ImageSettings())
        for values in [dict(output_format='jpeg', background='transparent'), dict(compression=10),
                       dict(size='1025x1024'), dict(size='4000x2160'), dict(size='1024x256'),
                       dict(size='512x512'), dict(size='3840x3840'), dict(size='bad'),
                       dict(output_format='webp', compression=101), dict(model='unverified')]:
            with self.subTest(values=values), self.assertRaises(ValueError):
                ImageSettings(**values)
        self.assertTrue(ImageSettings(size='3840x2160').experimental)
        self.assertFalse(ImageSettings(size='1536x1024').experimental)
        self.assertEqual(ImageSettings(output_format='webp', compression=75).api_params()['output_compression'], 75)
        self.assertNotIn('output_compression', ImageSettings().api_params())

    def test_output_formats_params_and_exclusive_creation(self):
        for fmt in ('png', 'jpeg', 'webp'):
            with self.subTest(fmt=fmt):
                settings = ImageSettings(model='gpt-image-2.5-flare', output_format=fmt,
                                         background='opaque' if fmt=='jpeg' else 'transparent', quality='xhigh',
                                         size='1536x864', compression=None if fmt=='png' else 25)
                self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('E2', 'output.'+fmt))
                config = load_config(self.queue)
                config['image_settings'] = settings
                job = make_job(load_queue(self.queue).iloc[0], config, True)
                generator = self.generator(settings)
                generator.generate(job)
                self.assertEqual(self.api.call_args.kwargs['size'], '1536x864')
                self.assertEqual(self.api.call_args.kwargs['background'], settings.background)
                self.assertEqual(self.api.call_args.kwargs['model'], settings.model)
                self.assertEqual(self.api.call_args.kwargs['quality'], 'xhigh')
                self.assertEqual(self.api.call_args.kwargs.get('output_compression'), settings.compression)
                self.assertTrue(verified_outputs(job.saidas, fmt))
                before = self.api.call_count
                with self.assertRaises(FileExistsError):
                    generator.generate(job)
                self.assertEqual(self.api.call_count, before)

    def test_output_wrong_content_is_uncertain_and_not_saved(self):
        settings = ImageSettings(output_format='jpeg')
        self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('E2', 'mismatch.jpg'))
        job = make_job(load_queue(self.queue).iloc[0], {**load_config(self.queue), 'image_settings': settings}, True)
        self.api.side_effect = None
        self.api.return_value = SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(png_bytes()).decode())])
        with self.assertRaises(GenerationError) as error:
            self.generator(settings).generate(job)
        self.assertTrue(error.exception.uncertain)
        self.assertFalse(job.saidas[0].exists())

    def test_changed_reference_after_preflight_does_not_call_api(self):
        job = make_job(load_queue(self.queue).iloc[0], load_config(self.queue), True)
        (self.refs / 'ok.png').write_bytes(image_bytes('png'))
        with self.assertRaisesRegex(GenerationError, 'alterada'):
            self.generator().generate(job)
        self.assertEqual(self.api.call_count, 0)

    def test_preferences_persist_only_whitelisted_nonsecret_data(self):
        values = deepcopy(self.session.preferences)
        values['image'] = asdict(ImageSettings(output_format='webp', background='transparent', compression=30))
        self.session.set_preferences(values, persist=True)
        again = StudioSession(self.session.settings_path, env={})
        self.assertEqual(again.preferences, values)
        self.assertEqual(again.photos, [])
        self.assertIsNone(again.queue)
        self.assertNotIn('OPENAI_API_KEY', self.session.settings_path.read_text())
        values['OPENAI_API_KEY'] = 'forbidden'
        with self.assertRaisesRegex(ValueError, 'credenciais'):
            self.session.set_preferences(values, persist=True)

    def test_simulation_pause_resume_and_mutation_lock(self):
        values = deepcopy(self.session.preferences)
        values.update(max_jobs=8, max_images=8, safe_mode=False)
        self.session.set_preferences(values)
        before = digest(self.queue)
        self.session.start()
        self.session.command('pause')
        count = len(self.session.state()['items'])
        time.sleep(.4)
        self.assertEqual(len(self.session.state()['items']), count)
        for mutation in [lambda: self.session.set_mode('direct'),
                         lambda: self.session.add_photos([self.refs/'ok.png']),
                         lambda: self.session.choose_queue(self.queue),
                         lambda: self.session.set_preferences(values),
                         lambda: self.session.set_folders([], '')]:
            with self.assertRaisesRegex(ValueError, 'ativa'):
                mutation()
        self.session.command('resume')
        self.session.worker.join(10)
        self.assertFalse(self.session.active)
        self.assertEqual(self.session.result['success'], 8)
        self.assertEqual(digest(self.queue), before)
        self.assertFalse(Path(str(self.queue)+'.state.json').exists())
        self.assertEqual(self.api.call_count, 0)
        log_path = Path(str(self.queue) + '.studio-run.json')
        log = json.loads(log_path.read_text(encoding='utf-8'))
        self.assertEqual(log['status'], 'Finalizado')
        self.assertEqual(log['result']['success'], 8)
        self.assertEqual(len([event for event in log['events'] if event['type'] == 'item_end']), 8)
        self.assertTrue(all(event['elapsed_seconds'] >= 0 for event in log['events'] if event['type'] == 'item_end'))
        self.assertNotIn('OPENAI_API_KEY', log_path.read_text(encoding='utf-8'))
        archive = Path(str(self.queue) + '.studio-runs') / (log['run_id'] + '.json')
        self.assertEqual(json.loads(archive.read_text(encoding='utf-8')), log)
        reopened = StudioSession(self.root / 'other-settings.json', env={})
        reopened.choose_queue(self.queue)
        self.assertEqual(reopened.state()['run_log']['result']['success'], 8)
        with self.assertRaisesRegex(ValueError, 'bloqueada'):
            self.session.start(real=True)

    def test_recovery_settings_refs_direct_ignored_column_and_history(self):
        overrides = {'reference_mode': 'direct', 'direct_references': (self.refs/'ok.png',)}
        main(self.queue, generator=self.generator(), env=self.env, overrides=overrides)
        self.assertEqual(self.api.call_count, 8)
        journal = Journal(self.queue)
        self.assertEqual(journal.records['1']['execution']['reference_mode'], 'direct')
        self.assertIn('sha256', journal.records['1']['execution']['references'][0])
        self.edit_queue(lambda b: (b['Fila_Geracao'].__setitem__('D2','INVALID::missing'), b['Fila_Geracao'].__setitem__('F2','PROCESSANDO')))
        main(self.queue, generator=self.generator(), env=self.env, overrides=overrides)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'CONCLUIDO')
        self.edit_queue(lambda b: b['Fila_Geracao'].__setitem__('F2','PROCESSANDO'))
        different = ImageSettings(background='opaque')
        main(self.queue, generator=self.generator(different), env=self.env, overrides={**overrides, 'image_settings':different})
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(self.api.call_count, 8)
        (self.refs/'ok.png').write_bytes(png_bytes())
        # Original generated bytes were also PNG but a different image: all completed histories block.
        (self.refs/'ok.png').write_bytes(image_bytes('png'))
        main(self.queue, generator=self.generator(), env=self.env, overrides=overrides)
        self.assertTrue(load_queue(self.queue)['Status'].eq('REVISAO').all())
        self.assertEqual(self.api.call_count, 8)

    def test_http_auth_origin_api_block_and_selected_files_only(self):
        server = make_server(self.session, picker=lambda kind: [str(self.queue)] if kind=='queue' else [])
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        def request(path, data=None, token=None, origin=None):
            connection = http.client.HTTPConnection('127.0.0.1', server.server_port)
            headers = {'Content-Type':'application/json'}
            if token:
                headers['X-Afirma-Token']=token
            if origin:
                headers['Origin']=origin
            connection.request('GET' if data is None else 'POST', path, body=None if data is None else json.dumps(data), headers=headers)
            response=connection.getresponse()
            body=response.read()
            connection.close()
            return response.status,body
        self.assertEqual(server.server_address[0], '127.0.0.1')
        self.assertEqual(request('/api/state')[0],403)
        self.assertEqual(request('/api/state',token=server.session_token,origin='https://untrusted.test')[0],403)
        self.assertEqual(request('/api/state',token=server.session_token)[0],200)
        status, body=request('/api/start',{'real':True},token=server.session_token)
        self.assertEqual(status,400)
        self.assertIn('bloqueada',body.decode())
        self.assertEqual(request('/api/thumbnail',{'index':0},token=server.session_token)[0],400)
        self.assertEqual(request('/.env',token=server.session_token)[0],404)
        dummy_key = 'sk-test-placeholder-only-1234567890'
        self.assertEqual(request('/api/credential', {'key': dummy_key}, origin='https://untrusted.test',
                                 token=server.session_token)[0], 403)
        status, body = request('/api/credential', {'key': dummy_key}, token=server.session_token)
        self.assertEqual(status, 200)
        self.assertNotIn(dummy_key.encode(), body)
        status, body = request('/api/state', token=server.session_token)
        self.assertEqual(json.loads(body)['api_key_status'], 'session')
        self.assertNotIn(dummy_key.encode(), body)
        self.assertEqual(request('/api/credential/clear', {}, token=server.session_token)[0], 200)
        self.assertEqual(self.session.api_key_status(), 'missing')

    def test_non_png_interrupted_recovery_uses_original_format(self):
        for fmt in ('jpeg', 'webp'):
            with self.subTest(fmt=fmt):
                def reset(book):
                    sheet=book['Fila_Geracao']
                    for i in range(2,10):
                        sheet.cell(i,5).value=f'{fmt}_{i}.{fmt}'
                        sheet.cell(i,6).value='PENDENTE'
                reset_queue = self.root/f'{fmt}.xlsx'
                reset_queue.write_bytes(self.queue.read_bytes())
                settings=ImageSettings(output_format=fmt)
                old_queue=self.queue
                self.queue=reset_queue
                try:
                    self.edit_queue(reset)
                    overrides={'reference_mode':'direct','direct_references':(self.refs/'ok.png',),'image_settings':settings,
                               'result_dir':self.root/f'{fmt}-results'}
                    main(self.queue,generator=self.generator(settings),env=self.env,overrides=overrides)
                    self.edit_queue(lambda b:b['Fila_Geracao'].__setitem__('F2','PROCESSANDO'))
                    before=self.api.call_count
                    # Session snapshot only includes pending jobs; recovery must not
                    # treat that pending subset as a reference mismatch.
                    overrides['expected_reference_hashes']={}
                    main(self.queue,generator=self.generator(settings),env=self.env,overrides=overrides)
                    self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'CONCLUIDO')
                    self.assertEqual(self.api.call_count,before)
                    self.assertEqual(Journal(self.queue).records['1']['execution']['settings']['output_format'],fmt)
                finally:
                    self.queue=old_queue

    def test_run_snapshot_rejects_reference_change_between_items(self):
        config,prefs=self.session.snapshot()
        config['expected_reference_hashes']={str(self.refs/'ok.png'):digest(self.refs/'ok.png')}
        (self.refs/'ok.png').write_bytes(image_bytes('png'))
        with self.assertRaisesRegex(ValueError,'durante a execução'):
            make_job(load_queue(self.queue).iloc[0],config,True)

    def test_direct_mode_ignores_unused_missing_folder_registry(self):
        def folders(book):
            sheet=book.create_sheet('Pastas_Referencias')
            sheet.append(['Alias','Caminho'])
            sheet.append(['Unused',str(self.root/'missing')])
        self.edit_queue(folders)
        self.session.set_mode('direct')
        self.session.add_photos([self.refs/'ok.png'])
        self.assertEqual(len(self.session.validate_campaign()['jobs']),8)

    def test_native_picker_contract_and_http_selector(self):
        from src.native_picker import select
        with patch('src.native_picker.tk.Tk') as root, patch('src.native_picker.filedialog.askopenfilenames',return_value=('a.png','b.png')):
            self.assertEqual(select('photos'),['a.png','b.png'])
            root.return_value.destroy.assert_called_once()
        with patch('src.native_picker.tk.Tk'), patch('src.native_picker.filedialog.askdirectory',return_value=str(self.refs)):
            self.assertEqual(select('folder'),[str(self.refs)])

    def test_http_queue_selection_returns_summary_to_fresh_session(self):
        fresh = StudioSession(self.root / "fresh-settings.json", env={})
        server = make_server(fresh, picker=lambda kind: [str(self.queue)] if kind == "queue" else [])
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)

        def request(path, data=None):
            connection = http.client.HTTPConnection("127.0.0.1", server.server_port)
            headers = {"X-Afirma-Token": server.session_token, "Content-Type": "application/json"}
            connection.request("GET" if data is None else "POST", path,
                               body=None if data is None else json.dumps(data), headers=headers)
            response = connection.getresponse()
            body = response.read()
            status = response.status
            connection.close()
            return status, json.loads(body)

        status, selected = request("/api/select", {"kind": "queue"})
        self.assertEqual(status, 200)
        self.assertEqual(selected["paths"], [str(self.queue)])
        status, state = request("/api/state")
        self.assertEqual(status, 200)
        self.assertEqual(state["queue"], str(self.queue))
        self.assertEqual(state["queue_summary"]["pending_items"], 8)
        self.assertEqual(state["queue_summary"]["execution_items"], 1)
        self.assertEqual(state["queue_summary"]["execution_images"], 1)
        self.assertEqual(self.api.call_count, 0)
    def test_session_folder_override_does_not_require_sheet_write(self):
        second=self.root/'new-photos'
        second.mkdir()
        (second/'ok.png').write_bytes(png_bytes())
        self.session.set_folders([{'alias':'Novas','path':str(second)}],'Novas')
        before=digest(self.queue)
        config,_=self.session.snapshot()
        result=main(self.queue,env={**self.env,'DRY_RUN':'SIM'},overrides=config)
        self.assertEqual(result['success'],8)
        self.assertEqual(digest(self.queue),before)
        self.assertEqual(make_job(load_queue(self.queue).iloc[0],config,True).referencias,(second/'ok.png',))

    def test_stop_reports_only_items_actually_processed(self):
        import contextlib
        from src.studio_session import RunControl
        control=RunControl(lambda *args:None)
        control.command('stop')
        output=io.StringIO()
        with contextlib.redirect_stdout(output):
            summary=main(self.queue,env={**self.env,'DRY_RUN':'SIM'},control=control)
        self.assertEqual(summary['processed'],0)
        self.assertEqual(summary['ignored'],8)
        self.assertIn('Processados: 0',output.getvalue())
        self.assertEqual(self.api.call_count,0)
