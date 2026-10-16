import base64
from copy import deepcopy
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock
from openpyxl import load_workbook
from main import main
from src.studio_session import StudioSession
from src.execution_state import Journal, digest
from src.excel_reader import load_queue
from src.image_generator import OpenAIImageGenerator
from support import IsolatedTest, png_bytes
import test_stage_three as fixtures


class OutputNamingFixture(IsolatedTest):
    def setUp(self):
        super().setUp()
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.queue, self.refs, self.output = fixtures.StageThreeTests().make_queue(self.root,
            [[i, 'Produto', '', 'ok.png', 'produto', 'PENDENTE', 0, ''] for i in (1,2)])
        self.session = StudioSession(self.root/'prefs.json', env={})
        self.session.choose_queue(self.queue)
        self.session.set_folders([dict(alias='Produtos',path=str(self.refs))], '')
        self.session.set_preferences({**self.session.preferences,'output_dir':str(self.output),
            'max_jobs':2,'max_images':2,'safe_mode':False})
        (self.output/'produto.png').write_bytes(png_bytes())
        self.env = dict(IMAGE_PROVIDER='openai',DRY_RUN='NAO',FIRST_RUN_SAFE_MODE='NAO',
                       MAX_JOBS_PER_RUN='2',MAX_IMAGES_PER_RUN='2')

    def generator(self, callback=None):
        def edit(**kw):
            if callback:
                callback()
            return SimpleNamespace(data=[SimpleNamespace(
                b64_json=base64.b64encode(png_bytes()).decode()) for _ in range(kw['n'])])
        self.api = Mock(side_effect=edit)
        return OpenAIImageGenerator(client=SimpleNamespace(images=SimpleNamespace(edit=self.api)),
                                    settings=self.session.snapshot()[0]['image_settings'])


class OutputNamingTests(OutputNamingFixture):
    def test_preview_is_read_only_and_numbers_disk_and_in_batch_collisions(self):
        before = digest(self.queue)
        plan = self.session.validate_campaign()
        self.assertEqual([Path(j['outputs'][0]).name for j in plan['jobs']],
                         ['produto_02.png','produto_03.png'])
        self.assertEqual(self.session.validate_campaign()['output_allocations'],plan['output_allocations'])
        self.assertEqual(digest(self.queue),before)
        self.assertFalse(Path(str(self.queue)+'.state.json').exists())

    def test_recorded_suffix_is_saved_before_call_and_recovered_without_new_call(self):
        def verify_before_call():
            records = Journal(self.queue).records
            live = [r for r in records.values() if r['phase']=='in_flight']
            self.assertEqual(len(live),1)
            self.assertEqual(Path(live[0]['outputs'][0]).name,
                             live[0]['output_allocation']['names'][0])
        generator = self.generator(verify_before_call)
        config,_ = self.session.snapshot()
        result = main(self.queue,env=self.env,generator=generator,overrides=config)
        self.assertEqual(result['success'],2)
        self.assertEqual(self.api.call_count,2)
        self.assertEqual((self.output/'produto.png').read_bytes(),png_bytes())
        self.assertEqual(load_queue(self.queue)['Nome_Saida'].tolist(),['produto','produto'])
        # Crash after files/history saved but before the spreadsheet completion.
        book=load_workbook(self.queue)
        book['Fila_Geracao']['F2']='PROCESSANDO'
        book.save(self.queue);book.close()
        restored,_=self.session.snapshot()
        main(self.queue,env=self.env,generator=generator,overrides=restored)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'CONCLUIDO')
        self.assertEqual(self.api.call_count,2)
        self.assertFalse((self.output/'produto_04.png').exists())

    def test_frozen_preview_conflict_never_renumbers_or_calls_api(self):
        config,_=self.session.snapshot()
        chosen=self.output/config['output_allocations']['1']['names'][0]
        chosen.write_bytes(b'external file')
        generator=self.generator()
        # Limit to the conflicting first item.
        result=main(self.queue,env={**self.env,'MAX_JOBS_PER_RUN':'1'},
                    generator=generator,overrides=config)
        self.assertEqual(result['errors'],1)
        self.api.assert_not_called()
        self.assertEqual(chosen.read_bytes(),b'external file')

    def test_multiple_images_are_allocated_as_one_group(self):
        book=load_workbook(self.queue)
        sheet=book['Fila_Geracao']
        sheet['I1']='Quantidade';sheet['I2']=2;sheet['I3']=1
        book.save(self.queue);book.close()
        (self.output/'produto_01.png').write_bytes(png_bytes())
        plan=self.session.validate_campaign()
        names=[Path(p).name for p in plan['jobs'][0]['outputs']]
        self.assertEqual(names,['produto_02_01.png','produto_02_02.png'])

    def test_recorded_paths_cannot_escape_output_directory(self):
        from main import make_job
        config,_=self.session.snapshot()
        config['output_allocations']['1']['names']=['../outside.png']
        with self.assertRaises(ValueError):
            make_job(load_queue(self.queue).iloc[0],config,True)

    def test_skips_existing_suffixes_and_keeps_selected_format(self):
        for extension in ('jpeg','webp'):
            with self.subTest(extension=extension):
                for suffix in ('','_02'):
                    (self.output/f'produto{suffix}.{extension}').write_bytes(b'preserved')
                prefs=deepcopy(self.session.preferences)
                prefs['image']['output_format']=extension
                self.session.set_preferences(prefs)
                plan=self.session.validate_campaign()
                self.assertEqual(Path(plan['jobs'][0]['outputs'][0]).name,f'produto_03.{extension}')

    def test_uncertain_history_keeps_its_suffix_and_never_becomes_new_pending(self):
        import httpx
        from openai import APITimeoutError
        generator=self.generator()
        self.api.side_effect=APITimeoutError(request=httpx.Request('POST','https://example.invalid'))
        generator.retries=0
        config,_=self.session.snapshot()
        main(self.queue,env={**self.env,'MAX_JOBS_PER_RUN':'1'},generator=generator,overrides=config)
        before=deepcopy(Journal(self.queue).records['1'])
        self.assertEqual(before['phase'],'uncertain')
        restored,_=self.session.snapshot()
        main(self.queue,env={**self.env,'MAX_JOBS_PER_RUN':'1'},generator=generator,overrides=restored)
        self.assertEqual(Journal(self.queue).records['1']['output_allocation'],before['output_allocation'])
        self.assertEqual(self.api.call_count,2)  # First item once, then only the other pending item.
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'REVISAO')
