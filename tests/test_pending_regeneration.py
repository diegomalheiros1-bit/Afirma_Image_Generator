"""Explicit sheet resets regenerate completed Studio items without losing evidence."""
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from main import main
from src.excel_reader import load_queue
from src.execution_state import Journal, PersistenceError, digest
from src.run_selection import plan_selection
from test_output_naming import OutputNamingFixture


class PendingRegenerationTests(OutputNamingFixture):
    # Reuse the isolated fixture, not the parent test cases in this module.
    def reset_sheet(self, statuses=('PENDENTE', 'PENDENTE'), change_prompt=False):
        book = load_workbook(self.queue)
        sheet = book['Fila_Geracao']
        for line, status in enumerate(statuses, 2):
            sheet.cell(line, 6).value = status
            if change_prompt:
                sheet.cell(line, 3).value = 'Nova composição com modelo'
        book.save(self.queue)
        book.close()

    def generate_first(self):
        generator = self.generator()
        config, _ = self.session.snapshot()
        self.assertEqual(main(self.queue, env=self.env, generator=generator,
                              overrides=config)['success'], 2)
        return generator

    def test_completed_pending_rows_preview_simulate_and_regenerate_with_audit(self):
        generator = self.generate_first()
        before = deepcopy(Journal(self.queue).records)
        old_files = {p: digest(p) for record in before.values() for p in record['outputs']}
        for changed in (False, True):
            with self.subTest(changed_prompt=changed):
                self.reset_sheet(change_prompt=changed)
                sheet_hash = digest(self.queue)
                journal_hash = digest(Journal(self.queue).path)
                plan = self.session.validate_campaign()
                self.assertEqual(plan['selection']['selected_ids'], ['1', '2'])
                self.assertEqual(plan['execution_preview']['images'], 2)
                self.assertTrue(all(d.get('regeneration') for d in plan['selection']['decisions']))
                self.assertTrue(all(not Path(p).exists() for job in plan['jobs'] for p in job['outputs']))
                self.session.start(real=False)
                self.session.worker.join(10)
                self.assertEqual(self.session.result['success'], 2)
                self.assertEqual(digest(self.queue), sheet_hash)
                self.assertEqual(digest(Journal(self.queue).path), journal_hash)
                config, _ = self.session.snapshot()
                result = main(self.queue, env=self.env, generator=generator, overrides=config)
                self.assertEqual(result['success'], 2)
                records = Journal(self.queue).records
                for key in ('1', '2'):
                    original = deepcopy(before[key])
                    history = original.pop('previous_generations', [])
                    self.assertEqual(records[key]['previous_generations'], history + [original])
                before = deepcopy(records)
                self.assertEqual({p: digest(p) for p in old_files}, old_files)
                self.assertEqual(load_queue(self.queue)['Status'].tolist(), ['CONCLUIDO'] * 2)
        self.assertEqual(self.api.call_count, 6)

    def test_five_rows_include_reset_completed_items_and_respect_limits(self):
        generator = self.generate_first()
        self.reset_sheet()
        book = load_workbook(self.queue)
        for key in (3, 4, 5):
            book['Fila_Geracao'].append([key, 'Produto', '', 'ok.png', 'produto', 'PENDENTE', 0, ''])
        book.save(self.queue)
        book.close()
        limited = self.session.validate_campaign()
        self.assertEqual(limited['selection']['selected_ids'], ['1', '2'])
        self.session.set_preferences({**self.session.preferences, 'max_jobs': 5, 'max_images': 5})
        full = self.session.validate_campaign()
        self.assertEqual(full['selection']['selected_ids'], ['1', '2', '3', '4', '5'])
        self.assertEqual(full['execution_preview']['images'], 5)
        self.assertEqual(len({p for job in full['jobs'] for p in job['outputs']}), 5)
        # Resume/errors do not turn an old completion into a new generation.
        self.session.set_selection('resume', ids=['1', '2', '3', '4', '5'])
        self.assertEqual(self.session.validate_campaign()['selection']['selected_ids'], ['3', '4', '5'])
        self.session.set_selection('errors')
        self.assertEqual(self.session.validate_campaign()['selection']['selected_ids'], [])
        self.session.set_selection('pending')
        config, _ = self.session.snapshot()
        result = main(self.queue, env={**self.env, 'MAX_JOBS_PER_RUN': '5', 'MAX_IMAGES_PER_RUN': '5'},
                      generator=generator, overrides=config)
        self.assertEqual(result['success'], 5)
        self.assertEqual(self.api.call_count, 7)
        self.assertEqual(load_queue(self.queue)['Status'].tolist(), ['CONCLUIDO'] * 5)
        self.assertEqual([len(r.get('previous_generations', []))
                          for r in Journal(self.queue).records.values()], [1, 1, 0, 0, 0])

    def test_crash_before_sheet_save_archives_once_and_resumes_the_same_new_name(self):
        generator = self.generate_first()
        self.reset_sheet(('PENDENTE', 'CONCLUIDO'))
        config, _ = self.session.snapshot()
        original = deepcopy(Journal(self.queue).records['1'])
        with patch('main.save_queue', side_effect=PersistenceError('disk unavailable')):
            with self.assertRaises(PersistenceError):
                main(self.queue, env=self.env, generator=generator, overrides=config)
        prepared = Journal(self.queue).records['1']
        self.assertEqual(prepared['phase'], 'prepared')
        self.assertEqual(prepared['previous_generations'], [original])
        self.assertEqual(self.api.call_count, 2)
        restored, _ = self.session.snapshot()
        result = main(self.queue, env=self.env, generator=generator, overrides=restored)
        self.assertEqual(result['success'], 1)
        record = Journal(self.queue).records['1']
        self.assertEqual(record['outputs'], prepared['outputs'])
        self.assertEqual(record['previous_generations'], [original])
        self.assertEqual(self.api.call_count, 3)

    def test_crash_after_new_files_ready_recovers_without_another_generation(self):
        from src.excel_reader import save_queue
        generator = self.generate_first()
        self.reset_sheet(('PENDENTE', 'CONCLUIDO'))
        config, _ = self.session.snapshot()

        def crash_after_completion(frame, queue):
            if frame['Status'].eq('CONCLUIDO').all():
                raise PersistenceError('completion save interrupted')
            save_queue(frame, queue)

        with patch('main.save_queue', side_effect=crash_after_completion):
            with self.assertRaises(PersistenceError):
                main(self.queue, env=self.env, generator=generator, overrides=config)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'PROCESSANDO')
        restored, _ = self.session.snapshot()
        result = main(self.queue, env=self.env, generator=generator, overrides=restored)
        self.assertEqual(result['success'], 0)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'CONCLUIDO')
        self.assertEqual(self.api.call_count, 3)
        self.assertEqual(len(Journal(self.queue).records['1']['previous_generations']), 1)

    def test_uncertain_generation_cannot_be_released_by_changing_sheet_to_pending(self):
        import httpx
        from openai import APITimeoutError
        generator = self.generate_first()
        self.reset_sheet(('PENDENTE', 'CONCLUIDO'))
        self.api.side_effect = APITimeoutError(request=httpx.Request('POST', 'https://example.invalid'))
        generator.retries = 0
        config, _ = self.session.snapshot()
        main(self.queue, env=self.env, generator=generator, overrides=config)
        self.reset_sheet(('PENDENTE', 'CONCLUIDO'))
        self.assertEqual(self.session.validate_campaign()['selection']['selected_ids'], [])
        restored, _ = self.session.snapshot()
        main(self.queue, env=self.env, generator=generator, overrides=restored)
        self.assertEqual(self.api.call_count, 3)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(len(Journal(self.queue).records['1']['previous_generations']), 1)

    def test_stale_preview_or_local_collision_does_not_destroy_completed_evidence(self):
        generator = self.generate_first()
        self.reset_sheet(('PENDENTE', 'CONCLUIDO'))
        config, _ = self.session.snapshot()
        journal = Journal(self.queue)
        original = deepcopy(journal.records['1'])
        output = self.output / config['output_allocations']['1']['names'][0]
        output.write_bytes(b'created by another program')
        result = main(self.queue, env=self.env, generator=generator, overrides=config)
        self.assertEqual(result['errors'], 1)
        self.assertEqual(Journal(self.queue).records['1'], original)
        self.assertEqual(self.api.call_count, 2)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'PENDENTE')
        config, _ = self.session.snapshot()
        journal.put('1', audit_note='updated after preview')
        with self.assertRaisesRegex(ValueError, 'histórico mudou'):
            main(self.queue, env=self.env, generator=generator, overrides=config)
        self.assertEqual(self.api.call_count, 2)

    def test_cli_keeps_existing_recovery_contract(self):
        self.generate_first()
        self.reset_sheet()
        config, _ = self.session.snapshot()
        frame = load_queue(self.queue)
        records = Journal(self.queue).records
        plain = {**config, 'studio': False}
        selection = plan_selection(frame, records, plain, 5, 5)
        self.assertEqual(selection['selected_ids'], [])
