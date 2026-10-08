"""Readiness regressions and installed SDK contract, entirely offline."""
import base64
import io
import json
from pathlib import Path
import tempfile
from unittest.mock import patch

import httpx
from openai import OpenAI
from openpyxl import load_workbook
from PIL import Image

from main import make_job
from src.execution_state import Journal, PersistenceError, digest
from src.image_generator import OpenAIImageGenerator, GenerationError
from src.image_settings import ImageSettings
from src.studio_session import StudioSession
from support import IsolatedTest
import test_stage_three


class ReadinessTests(IsolatedTest):
    def setUp(self):
        super().setUp()
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.queue, self.refs, self.output = test_stage_three.StageThreeTests().make_queue(self.root, [
            [1, 'Produto', 'Fundo claro', 'ok.png', 'result.png', 'PENDENTE', 0, '']])
        self.session = StudioSession(self.root / 'settings.json', env={})
        self.session.choose_queue(self.queue)

    def test_invalid_queue_selection_preserves_previous_campaign(self):
        previous = self.session.state()
        invalid = self.root / 'invalid.xlsx'
        invalid.write_bytes(self.queue.read_bytes())
        book = load_workbook(invalid)
        book['Fila_Geracao']['A1'] = 'invalid-column'
        book.save(invalid)
        book.close()
        with self.assertRaises(ValueError):
            self.session.choose_queue(invalid)
        self.assertEqual(self.session.state(), previous)

    def test_failed_mode_switch_preserves_direct_mode(self):
        self.session.set_mode('direct')
        book = load_workbook(self.queue)
        book['Fila_Geracao'].delete_cols(4)
        book.save(self.queue)
        book.close()
        with self.assertRaises(ValueError):
            self.session.set_mode('spreadsheet')
        self.assertEqual(self.session.mode, 'direct')

    def test_worker_refreshes_summary_after_persisted_completion(self):
        self.session.allow_api = True
        def completed(*args, **kwargs):
            book = load_workbook(self.queue)
            book['Fila_Geracao']['F2'] = 'CONCLUIDO'
            book.save(self.queue)
            book.close()
            return dict(processed=1, success=1, errors=0, ignored=0)
        with patch('main.main', side_effect=completed):
            self.session.start(real=True)
            self.session.worker.join(5)
        self.assertFalse(self.session.active)
        self.assertEqual(self.session.queue_summary['pending_items'], 0)
        self.assertEqual(self.session.queue_summary['completed_items'], 1)
        self.assertEqual(self.session.result['success'], 1)

    def test_malformed_journal_record_fails_closed(self):
        for record in (None, [], 'invalid', 42):
            with self.subTest(record=record):
                Path(str(self.queue) + '.state.json').write_text(json.dumps({'1': record}))
                with self.assertRaises(PersistenceError):
                    Journal(self.queue)

    def test_installed_sdk_multipart_formats_and_models(self):
        before = digest(self.queue)
        for model in ('gpt-image-2.5-sunburst', 'gpt-image-2.5-flare'):
            for fmt in ('png', 'jpeg', 'webp'):
                with self.subTest(model=model, fmt=fmt):
                    settings = ImageSettings(model=model, output_format=fmt,
                                             compression=None if fmt == 'png' else 80)
                    config, _ = self.session.snapshot()
                    config['image_settings'] = settings
                    from src.excel_reader import load_queue
                    row = load_queue(self.queue).iloc[0].copy()
                    row['Nome_Saida'] = f'{model}.{fmt}'
                    job = make_job(row, config, True)
                    self.output.mkdir(parents=True, exist_ok=True)
                    calls = []
                    def transport(request):
                        body = request.read()
                        calls.append(body)
                        self.assertEqual(request.url.path, '/v1/images/edits')
                        self.assertIn(b'multipart/form-data', request.headers['content-type'].encode())
                        self.assertIn(model.encode(), body)
                        self.assertIn(b'name="image[]"', body)
                        self.assertIn(b'name="output_format"', body)
                        self.assertEqual(b'name="output_compression"' in body, fmt != 'png')
                        buffer = io.BytesIO()
                        Image.new('RGB', (16, 16), 'red').save(buffer, format=fmt.upper())
                        return httpx.Response(200, json={'created': 1, 'data': [
                            {'b64_json': base64.b64encode(buffer.getvalue()).decode()}]})
                    with httpx.Client(transport=httpx.MockTransport(transport)) as http:
                        with OpenAI(api_key='offline-test-key', http_client=http, max_retries=0) as client:
                            generator = OpenAIImageGenerator(client=client, settings=settings, retries=0)
                            self.assertEqual(generator.generate(job), job.saidas)
                    self.assertEqual(len(calls), 1)
                    self.assertTrue(job.saidas[0].is_file())
        self.assertEqual(digest(self.queue), before)

    def test_installed_sdk_timeout_is_uncertain_without_hidden_retry(self):
        config, _ = self.session.snapshot()
        from src.excel_reader import load_queue
        job = make_job(load_queue(self.queue).iloc[0], config, True)
        calls = []
        def transport(request):
            calls.append(request.url.path)
            raise httpx.ReadTimeout('secret response must never escape', request=request)
        with httpx.Client(transport=httpx.MockTransport(transport)) as http:
            with OpenAI(api_key='offline-test-key', http_client=http, max_retries=0) as client:
                generator = OpenAIImageGenerator(client=client, retries=2)
                with self.assertRaises(GenerationError) as raised:
                    generator.generate(job)
        self.assertTrue(raised.exception.uncertain)
        self.assertNotIn('secret', str(raised.exception))
        self.assertEqual(len(calls), 1)
        self.assertFalse(job.saidas[0].exists())
