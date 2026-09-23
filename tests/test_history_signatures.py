import base64
import contextlib
import io
import tempfile
import hashlib
import json
import pandas as pd
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from main import main
from src.excel_reader import load_queue, save_queue
from src.execution_state import Journal, PersistenceError, row_fingerprint
from src.job_values import canonical_id, parse_quantity
from src.image_generator import OpenAIImageGenerator
from support import IsolatedTest, png_bytes
import test_stage_three


class HistoryTests(IsolatedTest):
    def setUp(self):
        super().setUp()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        rows = [self.row(i) for i in range(1, 4)]
        self.queue, self.refs, self.output = test_stage_three.StageThreeTests().make_queue(
            self.root, rows, optional=True)
        self.edit = Mock(side_effect=lambda **kw: SimpleNamespace(data=[SimpleNamespace(
            b64_json=base64.b64encode(png_bytes()).decode()) for _ in range(kw['n'])]))
        self.generator = OpenAIImageGenerator(client=SimpleNamespace(images=SimpleNamespace(edit=self.edit)), sleep=Mock())
        self.env = dict(IMAGE_PROVIDER="openai", DRY_RUN="NAO", FIRST_RUN_SAFE_MODE="NAO",
                        MAX_JOBS_PER_RUN="5", MAX_IMAGES_PER_RUN="5")

    def row(self, id, quantity=1):
        return [id, "base", "var", "ok.png", f"{id}.png", "PENDENTE", 0, None,
                "tema", "produto", "cliente", "negativo", quantity, "nota"]

    def run_queue(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return main(self.queue, generator=self.generator, env=self.env)

    def update(self, **cells):
        book = load_workbook(self.queue)
        for name, value in cells.items():
            book["Fila_Geracao"][name] = value
        book.save(self.queue)
        book.close()

    def append(self, id=4, quantity=None):
        book = load_workbook(self.queue)
        book["Fila_Geracao"].append(self.row(id, quantity))
        book.save(self.queue)
        book.close()

    def make_legacy(self, key='1'):
        journal = Journal(self.queue)
        for field in ("row_fingerprint", "row_fingerprint_version", "model", "quality", "result_dir"):
            journal.records[key].pop(field, None)
        journal.put(key)
        return journal

    @staticmethod
    def old_hash(row):
        # Frozen reproduction of the previous serialization, independent of v2.
        fields = ("ID", "Prompt_Padrao", "Prompt_Variacao", "Arquivo_Referencia", "Nome_Saida",
                  "Tema", "Produto", "Cliente", "Prompt_Negativo", "Quantidade", "Observacao_Usuario")
        pairs = []
        for field in fields:
            value = row.get(field, '')
            if value is None or isinstance(value, float) and value != value:
                value = ''
            pairs.append([field, str(value).strip()])
        return hashlib.sha256(json.dumps(pairs, ensure_ascii=True).encode()).hexdigest()

    def test_new_empty_quantity_keeps_completed_history(self):
        self.run_queue()
        book = load_workbook(self.queue)
        book["Fila_Geracao"].append(self.row(4, None))
        book.save(self.queue)
        book.close()
        self.run_queue()
        self.assertEqual(list(load_queue(self.queue)["Status"]), ["CONCLUIDO"] * 4)
        self.assertEqual(self.edit.call_count, 4)
        self.assertEqual(self.edit.call_args.kwargs['n'], 1)

    def test_legacy_without_signature_cannot_confirm_changed_prompt(self):
        self.run_queue()
        journal = Journal(self.queue)
        for field in ("row_fingerprint", "row_fingerprint_version", "model", "quality", "result_dir"):
            journal.records['1'].pop(field, None)
        journal.put('1')
        self.update(B2="changed prompt", F2="PROCESSANDO")
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(self.edit.call_count, 3)

    def test_equivalent_quantities_and_explicit_pandas_missing_values(self):
        expected = row_fingerprint({'ID': 1, 'Quantidade': 1})
        self.assertEqual(row_fingerprint({'ID': 1}), expected)
        for value in (None, '', ' ', 1, 1.0, '1', '1.0', float('nan'), pd.NA, pd.NaT):
            with self.subTest(value=repr(value)):
                self.assertEqual(parse_quantity(value), 1)
                self.assertEqual(row_fingerprint({'ID': 1, 'Quantidade': value}), expected)
        for value in (0, -1, 1.5, 'abc', '1.5', True, float('inf')):
            with self.subTest(invalid=repr(value)):
                with self.assertRaises(ValueError):
                    row_fingerprint({'ID': 1, 'Quantidade': value})
        self.assertNotEqual(expected, row_fingerprint({'ID': 1, 'Quantidade': 2}))

    def test_inferred_numeric_types_normalize_only_numeric_fields(self):
        integer = pd.DataFrame([{'ID': 1, 'Quantidade': 1}]).iloc[0]
        floated = pd.DataFrame([{'ID': 1, 'Quantidade': 1}, {'ID': 2, 'Quantidade': None}]).iloc[0]
        self.assertEqual(row_fingerprint(integer), row_fingerprint(floated))
        self.assertEqual(canonical_id(integer['ID']), canonical_id(floated['ID']))
        for field in ('Prompt_Padrao', 'Nome_Saida', 'Arquivo_Referencia', 'Cliente'):
            self.assertNotEqual(row_fingerprint({field: '001'}), row_fingerprint({field: '1'}))
        self.assertNotEqual(row_fingerprint({'ID': '001'}), row_fingerprint({'ID': '1'}))
        self.assertNotEqual(row_fingerprint({'ID': '1.0'}), row_fingerprint({'ID': 1}))

    def test_excel_preserves_text_ids_zeroes_and_na_text(self):
        self.update(A2='001', A3='NA', A4=3, B2='001')
        df = load_queue(self.queue)
        self.assertEqual(list(df['ID']), ['001', 'NA', 3])
        self.assertEqual(df.iloc[0]['Prompt_Padrao'], '001')
        self.run_queue()
        self.assertEqual(set(Journal(self.queue).records), {'001', 'NA', '3'})
        self.append('004')
        self.run_queue()
        self.assertEqual(self.edit.call_count, 4)
        self.assertTrue((load_queue(self.queue)['Status'] == 'CONCLUIDO').all())

    def test_unchanged_history_writes_neither_excel_nor_journal(self):
        self.run_queue()
        excel = self.queue.read_bytes()
        history = Journal(self.queue).path.read_bytes()
        with patch('main.save_queue', wraps=save_queue) as save, patch.object(Journal, 'put') as put:
            self.run_queue()
        save.assert_not_called()
        put.assert_not_called()
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(self.queue.read_bytes(), excel)
        self.assertEqual(Journal(self.queue).path.read_bytes(), history)

    def test_real_business_changes_still_block(self):
        self.run_queue()
        self.update(M2=2, B3='changed prompt', D4='other.png')
        self.run_queue()
        self.assertEqual(list(load_queue(self.queue)['Status']), ['REVISAO'] * 3)
        self.assertEqual(self.edit.call_count, 3)
        self.update(M2=1, B3='base', D4='ok.png', E2='renamed.png')
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(self.edit.call_count, 3)

    def test_v1_migration_requires_matching_historical_hash(self):
        self.run_queue()
        journal = Journal(self.queue)
        df = load_queue(self.queue)
        originals = {}
        for i, value in enumerate(('', 1, 1.0)):
            row = df.iloc[i].to_dict()
            row['Quantidade'] = value
            signature = self.old_hash(row)
            key = str(i + 1)
            originals[key] = signature
            journal.records[key]['row_fingerprint'] = signature
            journal.records[key].pop('row_fingerprint_version')
        journal.put('1')
        self.append(4, None)
        self.run_queue()
        self.assertEqual(self.edit.call_count, 4)
        self.assertTrue((load_queue(self.queue)['Status'] == 'CONCLUIDO').all())
        for key, old in originals.items():
            record = Journal(self.queue).records[key]
            self.assertEqual(record['row_fingerprint_version'], 2)
            self.assertEqual(record['row_fingerprint_v1'], old)

    def test_v1_mismatch_is_not_migrated(self):
        self.run_queue()
        journal = Journal(self.queue)
        old = self.old_hash(load_queue(self.queue).iloc[0])
        journal.records['1']['row_fingerprint'] = old
        journal.records['1'].pop('row_fingerprint_version')
        journal.put('1')
        original = journal.path.read_bytes()
        self.update(B2='changed', F2='PENDENTE')
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(journal.path.read_bytes(), original)
        self.assertEqual(self.edit.call_count, 3)

    def test_failed_signature_migration_blocks_new_task_and_preserves_history(self):
        self.run_queue()
        journal = Journal(self.queue)
        journal.records['1']['row_fingerprint'] = self.old_hash(load_queue(self.queue).iloc[0])
        journal.records['1'].pop('row_fingerprint_version')
        journal.put('1')
        original = journal.path.read_bytes()
        self.append(4)
        with patch('src.execution_state.os.replace', side_effect=PermissionError('locked')):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(journal.path.read_bytes(), original)
        self.assertEqual(load_queue(self.queue).iloc[3]['Status'], 'PENDENTE')

    def test_unknown_signature_version_blocks_recovery(self):
        self.run_queue()
        Journal(self.queue).put('1', row_fingerprint_version=999)
        self.update(F2='PROCESSANDO')
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertEqual(self.edit.call_count, 3)

    def test_legacy_completed_is_preserved_but_pending_or_processing_cannot_bypass(self):
        self.run_queue()
        journal = self.make_legacy()
        original = journal.path.read_bytes()
        with patch('main.save_queue', wraps=save_queue) as save:
            self.run_queue()
        save.assert_not_called()
        for status in ('PROCESSANDO', 'PENDENTE', 'REVISAO'):
            self.update(F2=status)
            self.run_queue()
            df = load_queue(self.queue)
            self.assertEqual(df.iloc[0]['Status'], 'REVISAO')
            self.assertIn('legado sem assinatura', df.iloc[0]['Observacao'])
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(journal.path.read_bytes(), original)

    def test_legacy_missing_invalid_or_modified_outputs_are_not_accepted(self):
        self.run_queue()
        for key in ('1', '2', '3'):
            self.make_legacy(key)
        (self.output / '1.png').unlink()
        (self.output / '2.png').write_bytes(b'invalid')
        with (self.output / '3.png').open('ab') as file:
            file.write(b'modified')
        self.run_queue()
        self.assertEqual(list(load_queue(self.queue)['Status']), ['REVISAO'] * 3)
        self.assertEqual(self.edit.call_count, 3)

    def test_new_record_recovers_with_changed_global_quality_and_no_charge(self):
        self.run_queue()
        self.update(F2='PROCESSANDO')
        self.generator.quality = 'high'
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'CONCLUIDO')
        self.assertEqual(Journal(self.queue).records['1']['quality'], 'auto')
        self.assertEqual(self.edit.call_count, 3)

    def test_failed_legacy_review_save_blocks_new_task(self):
        self.run_queue()
        self.make_legacy()
        self.update(F2='PROCESSANDO')
        self.append(4)
        with patch('main.save_queue', side_effect=PersistenceError('locked')):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.edit.call_count, 3)

    def test_legacy_key_ambiguity_cannot_create_second_history(self):
        self.run_queue()
        self.update(A2='001', F2='PENDENTE')
        original = Journal(self.queue).path.read_bytes()
        self.run_queue()
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'], 'REVISAO')
        self.assertNotIn('001', Journal(self.queue).records)
        self.assertEqual(Journal(self.queue).path.read_bytes(), original)
        self.assertEqual(self.edit.call_count, 3)
