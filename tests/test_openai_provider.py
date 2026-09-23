import base64
import contextlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from main import main
from src.file_manager import openai_output_paths
from src.generation_job import GenerationJob
from src.image_generator import OpenAIImageGenerator
import test_stage_three
from support import IsolatedTest, png_bytes
import httpx
from openai import RateLimitError, BadRequestError


class FakeImages:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def edit(self, **kwargs):
        self.calls.append(kwargs)
        result = next(self.responses)
        if isinstance(result, Exception):
            raise result
        return SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(x).decode())
                                     for x in result])


class OpenAITests(IsolatedTest):
    def test_first_real_run_preflight_blocks_unsafe_settings(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = [[1, "base", "var", "ok.png", "one.png", "PENDENTE", 0, None]]
            queue, refs, _ = test_stage_three.StageThreeTests().make_queue(root, rows)
            book = load_workbook(queue)
            book["Configuracao"]["B4"] = 1
            book.save(queue)
            book.close()
            safe = {"IMAGE_PROVIDER": "openai", "DRY_RUN": "NAO", "FIRST_RUN_SAFE_MODE": "SIM",
                    "MAX_JOBS_PER_RUN": "1", "MAX_IMAGES_PER_RUN": "1"}
            with patch.dict(os.environ, {**safe, "OPENAI_API_KEY": ""}):
                with self.assertRaisesRegex(ValueError, "OPENAI_API_KEY"):
                    main(queue)
            with patch.dict(os.environ, {**safe, "OPENAI_API_KEY": "test-only"}):
                with patch.dict(os.environ, {"MAX_JOBS_PER_RUN": "2"}):
                    with self.assertRaisesRegex(ValueError, "limites"):
                        main(queue)
                (refs / "ok.png").unlink()
                self.assertEqual(main(queue)["errors"], 1)
                (refs / "ok.png").touch()
                book = load_workbook(queue)
                book["Fila_Geracao"]["F2"] = "PENDENTE"
                book["Fila_Geracao"]["E2"] = "bad?.png"
                book.save(queue)
                book.close()
                self.assertEqual(main(queue)["errors"], 1)
            book = load_workbook(queue)
            self.assertEqual(book["Fila_Geracao"]["F2"].value, "ERRO")
            book.close()

    def test_api_inputs_retry_and_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            ref1, ref2 = root / "a.png", root / "b.jpg"
            ref1.write_bytes(png_bytes())
            ref2.write_bytes(png_bytes())
            paths = openai_output_paths(root, "imagem_001.png", 3)
            job = GenerationJob(1, "prompt\n\nEvitar: texto", (ref1, ref2), "imagem_001.png", "PROCESSANDO", 1,
                                quantidade=3, saidas=paths, prompt_negativo="texto")
            error = RateLimitError("secret", response=httpx.Response(429, request=httpx.Request("POST", "https://example.test")),
                                   body={"code": "rate_limit_exceeded"})
            api = FakeImages([error, [png_bytes()] * 3])
            generator = OpenAIImageGenerator(client=SimpleNamespace(images=api),
                                             sleep=lambda seconds: None, retries=1)
            self.assertEqual(generator.generate(job), paths)
            self.assertEqual([p.name for p in paths],
                             ["imagem_001_01.png", "imagem_001_02.png", "imagem_001_03.png"])
            self.assertEqual([p.read_bytes() for p in paths], [png_bytes()] * 3)
            self.assertEqual(len(api.calls), 2)
            self.assertEqual(api.calls[0]["n"], 3)
            self.assertEqual(api.calls[0]["output_format"], "png")
            self.assertEqual([Path(p.name).name for p in api.calls[0]["image"]], ["a.png", "b.jpg"])
            self.assertIn("Evitar: texto", api.calls[0]["prompt"])
            with self.assertRaises(FileExistsError):
                generator.generate(job)
            self.assertEqual(len(api.calls), 2)

    def test_real_flow_error_and_dry_run_no_call(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = [[1, "base", "var", "ok.png", "one.png", "PENDENTE", 0, None],
                    [2, "base", "var", "ok.png", "two.png", "PENDENTE", 0, None]]
            queue, _, _ = test_stage_three.StageThreeTests().make_queue(root, rows)
            class BadAPI:
                def __init__(self):
                    self.calls = 0
                def edit(self, **kwargs):
                    self.calls += 1
                    if self.calls == 1:
                        raise BadRequestError("sensitive payload", response=httpx.Response(400, request=httpx.Request("POST", "https://example.test")), body={})
                    return SimpleNamespace(data=[SimpleNamespace(
                        b64_json=base64.b64encode(png_bytes()).decode())])
            api = BadAPI()
            generator = OpenAIImageGenerator(client=SimpleNamespace(images=api))
            with patch.dict(os.environ, {"IMAGE_PROVIDER": "openai", "MAX_JOBS_PER_RUN": "2",
                                              "MAX_IMAGES_PER_RUN": "2"}):
                output = io.StringIO()
                with contextlib.redirect_stdout(output):
                    result = main(queue, generator=generator)
                self.assertIn("Provider: OpenAI", output.getvalue())
                self.assertEqual((result["success"], result["errors"]), (1, 1))
                book = load_workbook(queue)
                self.assertEqual(book["Fila_Geracao"]["F2"].value, "ERRO")
                self.assertEqual(book["Fila_Geracao"]["G2"].value, 1)
                self.assertNotIn("sensitive payload", book["Fila_Geracao"]["H2"].value)
                self.assertEqual(book["Fila_Geracao"]["F3"].value, "CONCLUIDO")
                book.close()
                self.assertEqual((root / "output" / "resultados" / "two.png").read_bytes(), png_bytes())

                book = load_workbook(queue)
                book["Configuracao"]["B5"] = "SIM"
                book["Fila_Geracao"]["F2"] = "PENDENTE"
                book.save(queue)
                book.close()
                calls = api.calls
                with contextlib.redirect_stdout(io.StringIO()):
                    main(queue, generator=generator)
                self.assertEqual(api.calls, calls)

    def test_safety_limits(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = [[i, "base", "var", "ok.png", f"{i}.png", "PENDENTE", 0, None,
                     None, None, None, None, 2, None] for i in range(1, 4)]
            queue, _, _ = test_stage_three.StageThreeTests().make_queue(root, rows, optional=True, dry_run="SIM")
            with patch.dict(os.environ, {"IMAGE_PROVIDER": "openai", "MAX_JOBS_PER_RUN": "3",
                                              "MAX_IMAGES_PER_RUN": "3"}):
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    result = main(queue)
                self.assertEqual(result["processed"], 1)
                self.assertIn("Imagens previstas: 2", out.getvalue())


if __name__ == "__main__":
    unittest.main()
