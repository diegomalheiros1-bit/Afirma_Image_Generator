import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main import main
from src.file_manager import validate_reference_paths
from src.image_generator import ImageGenerator, MockImageGenerator


class StageThreeTests(unittest.TestCase):
    def make_queue(self, root, rows, optional=False, dry_run="NAO"):
        references = root / "input" / "referencias"
        references.mkdir(parents=True)
        for name in ("modelo.jpg", "produto.webp", "ok.png"):
            (references / name).touch()
        output = root / "output" / "resultados"
        output.mkdir(parents=True)
        queue = root / "input" / "fila.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Fila_Geracao"
        columns = ["ID", "Prompt_Padrao", "Prompt_Variacao", "Arquivo_Referencia",
                   "Nome_Saida", "Status", "Tentativas", "Observacao"]
        if optional:
            columns += ["Tema", "Produto", "Cliente", "Prompt_Negativo", "Quantidade", "Observacao_Usuario"]
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        config = book.create_sheet("Configuracao")
        for pair in [("Parametro", "Valor"), ("Pasta_Referencias", "referencias"),
                     ("Pasta_Resultados", "resultados"), ("Limite_Por_Execucao", 20),
                     ("Dry_Run", dry_run)]:
            config.append(pair)
        book.save(queue)
        book.close()
        return queue, references, output

    def test_multiple_references_and_exact_missing_name(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            _, refs, _ = self.make_queue(root, [])
            self.assertEqual([p.name for p in validate_reference_paths(refs, " modelo.jpg ; produto.webp ")],
                             ["modelo.jpg", "produto.webp"])
            with self.assertRaisesRegex(FileNotFoundError, "missing.jpeg"):
                validate_reference_paths(refs, "ok.png;missing.jpeg")
            with self.assertRaisesRegex(ValueError, "bad.gif"):
                validate_reference_paths(refs, "ok.png;bad.gif")

    def test_mock_quantity_optional_columns_and_continuation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = [
                [1, "base", "var", "modelo.jpg;produto.webp", "first.png", "PENDENTE", 0, None,
                 "Mães", "Perfume", "Cliente", "sem texto", 3, "nota"],
                [2, "base", "var", "missing.jpeg", "second.png", "PENDENTE", 0, None,
                 None, None, None, None, 1, None],
                [3, "base", "var", "ok.png", "third.png", "PENDENTE", 0, None,
                 None, None, None, None, None, None],
            ]
            queue, _, output = self.make_queue(root, rows, optional=True)
            summary = main(queue, generator=MockImageGenerator())
            self.assertEqual((summary["success"], summary["errors"]), (2, 1))
            self.assertEqual(sorted(p.name for p in output.iterdir()),
                             ["first_001.png", "first_002.png", "first_003.png", "third.png"])
            book = load_workbook(queue)
            self.assertEqual([book["Fila_Geracao"][f"F{i}"].value for i in (2, 3, 4)],
                             ["CONCLUIDO", "ERRO", "CONCLUIDO"])
            self.assertIn("missing.jpeg", book["Fila_Geracao"]["H3"].value)
            book.close()

    def test_collision_and_mock_failure_do_not_overwrite(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = [[1, "base", "var", "ok.png", "first.png", "PENDENTE", 0, None,
                     None, None, None, None, 2, None],
                    [2, "base", "var", "ok.png", "next.png", "PENDENTE", 0, None,
                     None, None, None, None, 1, None]]
            queue, _, output = self.make_queue(root, rows, optional=True)
            (output / "first_002.png").write_bytes(b"original")
            summary = main(queue, generator=MockImageGenerator(fail=True))
            self.assertEqual(summary["errors"], 2)
            self.assertEqual((output / "first_002.png").read_bytes(), b"original")
            self.assertFalse((output / "first_001.png").exists())
            self.assertFalse((output / "next.png").exists())

    def test_abstract_interface(self):
        with self.assertRaises(TypeError):
            ImageGenerator()


if __name__ == "__main__":
    unittest.main()
