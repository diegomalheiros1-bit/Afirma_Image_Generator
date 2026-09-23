import base64
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch

import httpx
from openai import RateLimitError, APITimeoutError, AuthenticationError, InternalServerError
from openpyxl import load_workbook

from main import main
from src.excel_reader import load_queue, save_queue
from src.execution_state import Journal, PersistenceError, queue_lock
from src.image_generator import MockImageGenerator, OpenAIImageGenerator
from support import IsolatedTest, png_bytes
import test_stage_three


class ReliabilityTests(IsolatedTest):
    def setUp(self):
        super().setUp()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        rows = [[i, "base", "var", "ok.png", f"{i}.png", "PENDENTE", 0, None,
                 "tema", "produto", "cliente privado", "negativo", 1, "observação"] for i in range(1, 4)]
        self.queue, self.refs, self.output = test_stage_three.StageThreeTests().make_queue(
            self.root, rows, optional=True)
        self.env = {"IMAGE_PROVIDER": "openai", "DRY_RUN": "NAO", "FIRST_RUN_SAFE_MODE": "NAO",
                    "MAX_JOBS_PER_RUN": "3", "MAX_IMAGES_PER_RUN": "3"}
        self.calls = []
        self.edit = Mock(side_effect=self.respond)
        self.sleep = Mock()
        self.generator = OpenAIImageGenerator(client=SimpleNamespace(images=SimpleNamespace(edit=self.edit)),
                                             sleep=self.sleep)

    def respond(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(png_bytes()).decode())
                                     for _ in range(kwargs["n"])])

    def run_queue(self, **kwargs):
        with contextlib.redirect_stdout(io.StringIO()) as out:
            result = main(self.queue, generator=kwargs.pop("generator", self.generator),
                          env=kwargs.pop("env", self.env), **kwargs)
        self.console = out.getvalue()
        return result

    def update(self, **cells):
        book = load_workbook(self.queue)
        for key, value in cells.items():
            book["Fila_Geracao"][key] = value
        book.save(self.queue)
        book.close()

    def test_batch_prompt_and_configured_output(self):
        result = self.run_queue()
        self.assertEqual(result["success"], 3)
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(len(list(self.output.glob("*.png"))), 3)
        prompt = self.calls[0]["prompt"]
        for expected in ["base\n\nvar", "Tema: tema", "Produto: produto", "Evitar: negativo",
                         "Observação do usuário: observação"]:
            self.assertIn(expected, prompt)
        self.assertNotIn("cliente privado", prompt)
        records = Journal(self.queue).records
        self.assertEqual(records["1"]["api_attempts"], 1)
        self.assertEqual(records["1"]["task_attempts"], 1)
        self.assertNotIn("cliente privado", Journal(self.queue).path.read_text())

    def test_dry_run_matches_sent_prompt_and_changes_nothing(self):
        before = self.queue.read_bytes()
        self.run_queue(env={**self.env, "DRY_RUN": "SIM"})
        dry_output = self.console
        self.assertEqual(self.edit.call_count, 0)
        self.assertEqual(self.queue.read_bytes(), before)
        self.assertFalse(Journal(self.queue).path.exists())
        self.run_queue()
        self.assertIn(self.calls[0]["prompt"], dry_output)

    def test_mock_then_real(self):
        before = self.queue.read_bytes()
        self.run_queue(generator=MockImageGenerator())
        self.assertEqual(self.queue.read_bytes(), before)
        self.assertEqual(len(list(self.output.glob("*.png"))), 0)
        self.assertEqual(len(list((self.output / "_mock").rglob("*.png"))), 3)
        self.assertEqual(self.run_queue()["success"], 3)

    def test_limits_skip_large_job_and_select_later_small_job(self):
        self.update(M2=4)
        with self.assertLogs("gerador_imagens", "WARNING") as logs:
            result = self.run_queue(env={**self.env, "MAX_IMAGES_PER_RUN": "2"})
        self.assertEqual(result["processed"], 2)
        self.assertIn("Quantidade=4 excede", " ".join(logs.output))
        self.assertFalse((self.output / "1_01.png").exists())
        self.assertEqual(self.edit.call_count, 2)

    def test_first_test_protection_is_explicit_even_with_fake_client(self):
        with self.assertRaisesRegex(ValueError, "FIRST_RUN_SAFE_MODE"):
            self.run_queue(env={**self.env, "FIRST_RUN_SAFE_MODE": "SIM"})
        self.assertEqual(self.edit.call_count, 0)

    def test_excel_and_environment_job_caps(self):
        book = load_workbook(self.queue)
        book["Configuracao"]["B4"] = 1
        book.save(self.queue)
        book.close()
        self.assertEqual(self.run_queue()["processed"], 1)
        book = load_workbook(self.queue)
        book["Configuracao"]["B4"] = 20
        book.save(self.queue)
        book.close()
        self.assertEqual(self.run_queue(env={**self.env, "MAX_JOBS_PER_RUN": "1"})["processed"], 1)
        self.assertEqual(self.edit.call_count, 2)

    def test_excel_atomic_failure_before_call_and_recovery(self):
        original = self.queue.read_bytes()
        replace_file = os.replace
        def denied(source, target):
            if Path(target) == self.queue:
                raise PermissionError("Excel locked")
            return replace_file(source, target)
        with patch("src.excel_reader.os.replace", side_effect=denied):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.queue.read_bytes(), original)
        self.assertEqual(self.edit.call_count, 0)
        self.assertEqual(Journal(self.queue).records["1"]["phase"], "prepared")
        self.assertEqual(self.run_queue()["success"], 3)

    def test_excel_failure_after_generation_recovers_without_second_charge(self):
        calls = 0
        def persist(df, queue):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise PersistenceError("Excel locked at completion")
            save_queue(df, queue)
        with patch("main.save_queue", side_effect=persist):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.edit.call_count, 1)
        self.assertEqual(Journal(self.queue).records["1"]["phase"], "files_ready")
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "PROCESSANDO")
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.assertTrue((load_queue(self.queue)["Status"] == "CONCLUIDO").all())

    def test_interruption_never_replays_unknown_request_even_after_manual_pending(self):
        self.edit.side_effect = KeyboardInterrupt
        with self.assertRaises(KeyboardInterrupt):
            self.run_queue()
        self.assertEqual(Journal(self.queue).records["1"]["phase"], "in_flight")
        self.update(F2="PENDENTE")
        self.edit.side_effect = self.respond
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)  # interrupted call + jobs 2 and 3
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "REVISAO")

    def test_modified_file_never_regenerates(self):
        self.run_queue()
        (self.output / "1.png").write_bytes(b"changed")
        self.update(F2="PENDENTE")
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "REVISAO")

    def test_unregistered_processing_and_existing_output_block_replay(self):
        self.update(F2="PROCESSANDO")
        (self.output / "2.png").write_bytes(png_bytes())
        self.run_queue()
        self.assertEqual(self.edit.call_count, 1)
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "REVISAO")

    def test_journal_failure_before_api_stops_queue(self):
        original = Journal.put
        def fail(journal, key, **values):
            if values.get("phase") == "in_flight":
                raise PersistenceError("Disk full")
            original(journal, key, **values)
        with patch.object(Journal, "put", new=fail):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.edit.call_count, 0)

    def test_journal_failure_after_outputs_blocks_replay(self):
        original = Journal.put
        def fail(journal, key, **values):
            if values.get("phase") == "files_ready":
                raise PersistenceError("Disk full")
            original(journal, key, **values)
        with patch.object(Journal, "put", new=fail):
            with self.assertRaises(PersistenceError):
                self.run_queue()
        self.assertEqual(self.edit.call_count, 1)
        self.assertTrue((self.output / "1.png").exists())
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "REVISAO")

    def test_partial_output_is_preserved_and_not_completed(self):
        self.update(M2=2)
        from pathlib import Path
        original = Path.open
        def fail(path, mode="r", *args, **kwargs):
            if path.name == "1_02.png" and mode == "xb":
                raise OSError("Disk full")
            return original(path, mode, *args, **kwargs)
        with patch.object(Path, "open", new=fail):
            self.run_queue(env={**self.env, "MAX_IMAGES_PER_RUN": "4"})
        self.assertTrue((self.output / "1_01.png").exists())
        self.assertFalse((self.output / "1_02.png").exists())
        self.assertEqual(load_queue(self.queue).iloc[0]["Status"], "REVISAO")
        calls = self.edit.call_count
        self.run_queue()
        self.assertEqual(self.edit.call_count, calls)

    def test_os_lock_blocks_another_process_and_releases(self):
        code = ("from pathlib import Path; from src.execution_state import queue_lock, PersistenceError; "
                "import sys\ntry:\n with queue_lock(Path(sys.argv[1])): pass\n"
                "except PersistenceError: sys.exit(7)")
        with queue_lock(self.queue):
            result = subprocess.run([sys.executable, "-c", code, str(self.queue)], capture_output=True)
            self.assertEqual(result.returncode, 7, result.stderr)
        with queue_lock(self.queue):
            pass

    def api_error(self, cls, status, code, headers=None):
        return cls("secret-and-binary-should-not-leak", response=httpx.Response(status,
                   request=httpx.Request("POST", "https://example.test"), headers=headers or {}), body={"code": code})

    def test_retry_after_backoff_and_separate_counters(self):
        responses = [self.api_error(RateLimitError, 429, "slow_down", {"retry-after": "5"}),
                     self.api_error(InternalServerError, 503, "server_is_overloaded")]
        def edit(**kwargs):
            if responses:
                raise responses.pop(0)
            return self.respond(**kwargs)
        self.edit.side_effect = edit
        self.run_queue()
        self.assertEqual(self.edit.call_count, 5)
        self.assertGreaterEqual(self.sleep.call_args_list[0].args[0], 5)
        self.assertGreaterEqual(self.sleep.call_args_list[1].args[0], 2)
        record = Journal(self.queue).records["1"]
        self.assertEqual((record["task_attempts"], record["api_attempts"]), (1, 3))

    def test_quota_and_auth_errors_are_not_retried_or_logged_raw(self):
        self.edit.side_effect = self.api_error(RateLimitError, 429, "insufficient_quota")
        with self.assertLogs("gerador_imagens", "INFO") as logs:
            self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.sleep.assert_not_called()
        self.assertNotIn("secret-and-binary", " ".join(logs.output))
        self.assertTrue((load_queue(self.queue)["Status"] == "ERRO").all())
        self.update(F2="PENDENTE")
        self.edit.side_effect = self.api_error(AuthenticationError, 401, "invalid_api_key")
        self.run_queue()
        self.assertEqual(self.edit.call_count, 4)
        self.sleep.assert_not_called()

    def test_invalid_api_image_never_marks_completed(self):
        self.edit.side_effect = lambda **kwargs: SimpleNamespace(data=[SimpleNamespace(
            b64_json=base64.b64encode(b"not an image").decode())])
        self.run_queue()
        self.assertTrue((load_queue(self.queue)["Status"] == "REVISAO").all())
        self.assertFalse(list(self.output.glob("*.png")))

    def test_timeout_is_uncertain_and_not_retried(self):
        self.edit.side_effect = APITimeoutError(request=httpx.Request("POST", "https://example.test"))
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)
        self.sleep.assert_not_called()
        self.assertTrue((load_queue(self.queue)["Status"] == "REVISAO").all())
        self.run_queue()
        self.assertEqual(self.edit.call_count, 3)

    def test_exhausted_retries_and_long_retry_after(self):
        self.edit.side_effect = self.api_error(RateLimitError, 429, "slow_down")
        self.run_queue()
        self.assertEqual(self.edit.call_count, 9)  # 3 jobs, initial + 2 retries
        self.assertEqual(self.sleep.call_count, 6)
        self.update(F2="PENDENTE")
        self.edit.side_effect = self.api_error(RateLimitError, 429, "slow_down", {"retry-after": "600"})
        self.run_queue()
        self.assertEqual(self.edit.call_count, 10)
        self.assertEqual(self.sleep.call_count, 6)
