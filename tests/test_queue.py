import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main import is_pending, main
from src.file_manager import validate_reference_path
from src.prompt_builder import build_prompt


class QueueTests(unittest.TestCase):
    def test_prompt(self):
        self.assertEqual(build_prompt(" base ", " variante "), "base\n\nvariante")
        with self.assertRaises(ValueError):
            build_prompt(" ", "")

    def test_reference(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "photo.png").touch()
            self.assertEqual(validate_reference_path(root, "photo.png"), root / "photo.png")
            with self.assertRaises(FileNotFoundError):
                validate_reference_path(root, "missing.png")
            with self.assertRaises(ValueError):
                validate_reference_path(root, "../outside.png")

    def test_status_filter(self):
        self.assertTrue(is_pending(" pendente "))
        for status in ("CONCLUIDO", "ERRO", "PROCESSANDO", None):
            self.assertFalse(is_pending(status))

    def test_error_continues_and_preserves_configuration(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            references = root / "references"
            references.mkdir()
            (references / "ok.png").touch()
            queue = root / "queue.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Fila_Geracao"
            sheet.append(["ID", "Prompt_Padrao", "Prompt_Variacao", "Arquivo_Referencia",
                          "Nome_Saida", "Status", "Tentativas", "Observacao"])
            sheet.append([1, "base", "var", "missing.png", "a.png", "PENDENTE", 0, None])
            sheet.append([2, "base", "var", "ok.png", "b.png", "PENDENTE", 0, None])
            sheet.append([3, "base", "var", "ok.png", "c.png", "CONCLUIDO", 1, "anterior"])
            config = workbook.create_sheet("Configuracao")
            for entry in [("Parametro", "Valor"), ("Pasta_Referencias", str(references)),
                          ("Pasta_Resultados", str(root / "results")),
                          ("Limite_Por_Execucao", 10), ("Dry_Run", "NAO")]:
                config.append(entry)
            workbook.save(queue)
            workbook.close()

            main(queue, references)
            result = load_workbook(queue)
            self.assertEqual(result.sheetnames, ["Fila_Geracao", "Configuracao"])
            self.assertEqual(result["Configuracao"]["A1"].value, "Parametro")
            rows = result["Fila_Geracao"]
            self.assertEqual((rows["F2"].value, rows["G2"].value), ("ERRO", 1))
            self.assertEqual((rows["F3"].value, rows["G3"].value), ("CONCLUIDO", 1))
            self.assertEqual((rows["F4"].value, rows["G4"].value), ("CONCLUIDO", 1))
            main(queue, references)
            self.assertEqual(load_workbook(queue)["Fila_Geracao"]["G3"].value, 1)
            result.close()


if __name__ == "__main__":
    unittest.main()
