"""Local fixtures only. No credentials, SDK client or network permitted."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
from unittest.mock import patch
from openpyxl import load_workbook
from support import IsolatedTest, png_bytes
import test_stage_three as fixtures
from src.studio_session import StudioSession
from src.execution_state import Journal, execution_identity, row_fingerprint, verified_outputs
from src.excel_reader import load_queue
from src.run_selection import plan_selection, review_evidence


class ExecutionExperienceTests(IsolatedTest):
    def setUp(self):
        super().setUp()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.queue, refs, output = fixtures.StageThreeTests().make_queue(self.root,
            [[i, 'Produto', '', 'ok.png', str(i), 'PENDENTE', 0, ''] for i in range(1, 13)])
        self.session = StudioSession(self.root/'settings.json', allow_api=False, env={})
        self.session.choose_queue(self.queue)
        self.session.set_folders([{'alias': 'Produtos', 'path': str(refs)}], 'Produtos')
        self.session.set_preferences({**self.session.preferences, 'output_dir': str(output), 'max_jobs': 10, 'max_images': 10, 'safe_mode': False})
        self.config, self.prefs = self.session.snapshot()

    def selection(self, mode, records=None, ids=None, items=10, images=10):
        return plan_selection(load_queue(self.queue), records or {}, self.config, items, images, mode, ids)

    def set_statuses(self, statuses):
        book = load_workbook(self.queue)
        for i, status in enumerate(statuses, 2):
            book['Fila_Geracao'][f'F{i}'] = status
        book.save(self.queue)
        book.close()

    def prepare_record(self, key='1', phase='uncertain'):
        from main import make_job
        frame = load_queue(self.queue)
        row = frame[frame['ID'].map(str) == key].iloc[0]
        job = make_job(row, self.config, True)
        identity, signature = execution_identity(job, row)
        Journal(self.queue).put(key, phase=phase, row_fingerprint_version=2,
            row_fingerprint=row_fingerprint(row), execution=identity, execution_fingerprint=signature,
            outputs=[str(p) for p in job.saidas], api_attempts=1, task_attempts=1,
            error={'category':'connection','exception_type':'APIConnectionError','uncertain':True})
        return job

    def test_modes_exclude_completed_uncertain_and_unchanged_errors(self):
        self.set_statuses(['CONCLUIDO','REVISAO','PROCESSANDO','ERRO','ERRO','PENDENTE'])
        records = {'4': {'phase':'rejected','error':{'category':'authentication','uncertain':False}},
                   '5': {'phase':'rejected','error':{'category':'limit_quota','api_code':'rate_limit_exceeded','uncertain':False}}}
        self.assertEqual(self.selection('errors',records)['selected_ids'], ['5'])
        self.assertEqual(self.selection('pending',records)['selected_ids'], [str(i) for i in range(6,13)])
        self.assertEqual(self.selection('pending_errors',records)['selected_ids'], ['5']+[str(i) for i in range(6,13)])
        self.assertEqual(self.selection('resume',records,['1','2','3','5','6'])['selected_ids'], ['5','6'])

    def test_original_ten_only_three_unstarted_never_external(self):
        self.set_statuses(['CONCLUIDO']*6+['REVISAO']+['PENDENTE']*5)
        plan=self.selection('resume',{},[str(i) for i in range(1,11)])
        self.assertEqual(plan['selected_ids'], ['8','9','10'])
        self.assertEqual([d['category'] for d in plan['decisions'][-2:]],['outside','outside'])

    def test_caps_and_uncertain_pending_journal(self):
        records={'1':{'phase':'uncertain'}}
        plan=self.selection('pending',records,items=4,images=2)
        self.assertEqual(plan['selected_ids'], ['2','3'])
        self.assertEqual(plan['images'],2)
        self.assertEqual(plan['decisions'][0]['category'],'review')
        self.assertEqual(plan['decisions'][3]['category'],'limit')

    def test_local_rejection_only_eligible_after_validation_succeeds(self):
        self.set_statuses(['ERRO'])
        records={'1':{'phase':'local_rejected','local_validation_calls':0}}
        self.assertEqual(self.selection('errors',records)['selected_ids'],['1'])
        book=load_workbook(self.queue)
        book['Fila_Geracao']['D2']='missing.png'
        book.save(self.queue);book.close()
        self.assertEqual(self.selection('errors',records)['selected_ids'],[])

    def test_image_cap_counts_quantity_and_defers_whole_item(self):
        book=load_workbook(self.queue)
        book['Fila_Geracao']['I1']='Quantidade'
        book['Fila_Geracao']['I2']=3
        book.save(self.queue);book.close()
        plan=self.selection('pending',items=10,images=2)
        self.assertEqual(plan['selected_ids'],['2','3'])
        self.assertEqual(plan['decisions'][0]['category'],'limit')

    def moderation_record(self):
        from main import make_job
        from src.execution_state import fingerprint
        self.set_statuses(['ERRO'])
        row = load_queue(self.queue).iloc[0]
        job = make_job(row, self.config, True)
        identity, _ = execution_identity(job, row)
        return {'1': dict(phase='rejected', execution=identity,
            fingerprint=fingerprint(job, job.settings.model, job.settings.quality),
            error=dict(category='parameters', api_code='moderation_blocked', uncertain=False))}

    def test_moderation_retry_allows_unchanged_and_changed_content_without_mutation(self):
        from PIL import Image
        records = self.moderation_record()
        original = deepcopy(records)
        self.assertEqual(self.selection('errors', records)['selected_ids'], ['1'])
        self.assertNotIn('1', self.selection('pending', records)['selected_ids'])
        book = load_workbook(self.queue)
        book['Fila_Geracao']['C2'] = 'Modelo adulto em publicidade de cosméticos, com roupa.'
        book.save(self.queue); book.close()
        self.assertEqual(self.selection('errors', records)['selected_ids'], ['1'])
        self.assertEqual(records, original)
        # A change to the actual image bytes is also recognized with the same path.
        records = self.moderation_record()
        Image.new('RGB', (12, 12), 'green').save(self.queue.parent/'referencias'/'ok.png')
        self.assertEqual(self.selection('errors', records)['selected_ids'], ['1'])

    def test_moderation_retry_does_not_require_settings_output_or_metadata_edits(self):
        from dataclasses import replace
        records = self.moderation_record()
        self.config['image_settings'] = replace(self.config['image_settings'], quality='high')
        self.assertEqual(self.selection('errors', records)['selected_ids'], ['1'])
        book = load_workbook(self.queue)
        book['Fila_Geracao']['E2'] = 'different-output'
        book['Fila_Geracao']['H2'] = 'Apenas observação operacional'
        book.save(self.queue); book.close()
        self.assertEqual(self.selection('errors', records)['selected_ids'], ['1'])

    def test_moderation_retry_keeps_uncertain_missing_evidence_and_outputs_blocked(self):
        records = self.moderation_record()
        book = load_workbook(self.queue)
        book['Fila_Geracao']['C2'] = 'Prompt corrigido'
        book.save(self.queue); book.close()
        for mutation in ('uncertain', 'missing', 'output'):
            with self.subTest(mutation=mutation):
                candidate = deepcopy(records)
                if mutation == 'uncertain': candidate['1']['error']['uncertain'] = True
                if mutation == 'missing': candidate['1']['error'].pop('uncertain')
                if mutation == 'output': Path(candidate['1']['execution']['outputs'][0]).write_bytes(png_bytes())
                self.assertEqual(self.selection('errors', candidate)['selected_ids'], [])

    def test_old_history_requires_explicit_selection(self):
        self.session.run_log={'version':1,'planned_items':10,'events':[]}
        self.session.set_selection('resume')
        with self.assertRaisesRegex(ValueError,'IDs'):
            self.session.validate_campaign()
        self.session.set_selection('resume',ids=['8','9','10'])
        self.assertEqual(self.session.validate_campaign()['selection']['selected_ids'],['8','9','10'])

    def test_review_requires_consent_and_proof_then_preserves_audit(self):
        self.set_statuses(['REVISAO'])
        self.prepare_record()
        with self.assertRaises(ValueError):
            self.session.review_action('1','release')
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'REVISAO')
        self.session.review_action('1','release',True)
        record=Journal(self.queue).records['1']
        self.assertEqual(record['phase'],'retry_authorized')
        self.assertEqual(record['review_actions'][0]['prior_phase'],'uncertain')
        self.assertTrue(record['review_actions'][0]['duplicate_charge_acknowledged'])
        self.assertNotIn('1',self.selection('errors',{'1':record})['selected_ids'])
        self.assertNotIn('1',self.selection('pending',{'1':record})['selected_ids'])
        self.assertIn('1',self.selection('review',{'1':record})['selected_ids'])

    def test_review_mode_only_selects_explicitly_released_reviews(self):
        self.set_statuses(['REVISAO','REVISAO','ERRO','PENDENTE'])
        self.prepare_record('1')
        self.prepare_record('2')
        Journal(self.queue).put('3',phase='rejected',error=dict(category='limit_quota',api_code='rate_limit_exceeded',uncertain=False))
        self.assertEqual(self.selection('review',Journal(self.queue).records)['selected_ids'],[])
        self.session.review_action('1','release',True)
        records=Journal(self.queue).records
        self.assertEqual(self.selection('review',records)['selected_ids'],['1'])
        visible=self.session.reviews()
        self.assertTrue(next(item for item in visible if item['id']=='1')['released'])
        self.assertFalse(next(item for item in visible if item['id']=='2')['released'])
        self.assertEqual(load_queue(self.queue).iloc[1]['Status'],'REVISAO')
        self.assertFalse(self.session.active)

    def test_review_mode_revalidates_released_identity_and_caps(self):
        self.set_statuses(['REVISAO','REVISAO'])
        self.prepare_record('1');self.prepare_record('2')
        self.session.review_action('1','release',True)
        self.session.review_action('2','release',True)
        plan=self.selection('review',Journal(self.queue).records,items=1,images=1)
        self.assertEqual(plan['selected_ids'],['1'])
        self.assertEqual(plan['decisions'][1]['category'],'limit')
        book=load_workbook(self.queue);book['Fila_Geracao']['B2']='Changed';book.save(self.queue);book.close()
        self.assertEqual(self.selection('review',Journal(self.queue).records)['selected_ids'],['2'])

    def test_authorized_reviews_are_selected_even_if_sheet_still_says_review(self):
        self.set_statuses(['REVISAO']*4)
        for key in ['1','2','3']:
            self.prepare_record(key)
            self.session.review_action(key,'release',True)
        self.set_statuses(['REVISAO']*4)
        before=self.queue.read_bytes()
        self.session.set_selection('review')
        plan=self.session.validate_campaign()
        self.assertEqual(plan['selection']['selected_ids'],['1','2','3'])
        self.assertEqual((plan['execution_preview']['items'],plan['execution_preview']['images']),(3,3))
        self.assertEqual(self.queue.read_bytes(),before)
        self.session.start(real=False)
        self.session.worker.join(10)
        self.assertEqual(self.session.result['success'],3)
        self.assertEqual(self.queue.read_bytes(),before)
        self.assertFalse(self.session.allow_api)

    def test_authorized_review_with_divergence_or_existing_outputs_stays_excluded(self):
        self.set_statuses(['REVISAO','REVISAO'])
        first=self.prepare_record('1');self.prepare_record('2')
        self.session.review_action('1','release',True);self.session.review_action('2','release',True)
        self.set_statuses(['REVISAO','REVISAO'])
        first.saidas[0].write_bytes(png_bytes())
        book=load_workbook(self.queue);book['Fila_Geracao']['B3']='Changed';book.save(self.queue);book.close()
        self.assertEqual(self.selection('review',Journal(self.queue).records)['selected_ids'],[])

    def test_fake_generation_executes_authorized_review_labels_only(self):
        import base64
        from types import SimpleNamespace
        from unittest.mock import Mock
        from main import main
        from src.image_generator import OpenAIImageGenerator
        self.set_statuses(['REVISAO']*3)
        for key in ['1','2']:
            self.prepare_record(key);self.session.review_action(key,'release',True)
        self.set_statuses(['REVISAO']*3)
        edit=Mock(return_value=SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(png_bytes()).decode())]))
        generator=OpenAIImageGenerator(client=SimpleNamespace(images=SimpleNamespace(edit=edit)),retries=0)
        config={**self.config,'selection_mode':'review','dry_run':False}
        result=main(self.queue,generator=generator,overrides=config,env=dict(FIRST_RUN_SAFE_MODE='NAO',MAX_JOBS_PER_RUN='10',MAX_IMAGES_PER_RUN='10'))
        self.assertEqual(result['success'],2)
        self.assertEqual(edit.call_count,2)
        self.assertEqual(load_queue(self.queue)['Status'].tolist()[:3],['CONCLUIDO','CONCLUIDO','REVISAO'])

    def test_review_divergence_and_existing_output_forbid_release(self):
        self.set_statuses(['REVISAO'])
        job=self.prepare_record()
        job.saidas[0].write_bytes(png_bytes())
        with self.assertRaises(ValueError):
            self.session.review_action('1','release',True)
        record=Journal(self.queue).records['1']
        record['execution_fingerprint']='changed'
        self.assertFalse(review_evidence(load_queue(self.queue).iloc[0],record,self.config)['matching'])
        self.assertFalse(review_evidence(load_queue(self.queue).iloc[0],record,self.config)['recoverable'])

    def test_recover_only_verified_hashes_identity_configuration(self):
        self.set_statuses(['REVISAO'])
        job=self.prepare_record(phase='files_ready')
        job.saidas[0].write_bytes(png_bytes())
        Journal(self.queue).put('1',hashes=verified_outputs(job.saidas))
        self.assertTrue(self.session.reviews()[0]['evidence']['recoverable'])
        self.session.review_action('1','recover',True)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'CONCLUIDO')
        self.assertEqual(job.saidas[0].read_bytes(),png_bytes())

    def test_wrong_output_paths_cannot_recover_or_release_even_with_valid_hash(self):
        self.set_statuses(['REVISAO'])
        self.prepare_record(phase='files_ready')
        other=self.root/'unrelated.png';other.write_bytes(png_bytes())
        Journal(self.queue).put('1',outputs=[str(other)],hashes=verified_outputs([other]))
        evidence=self.session.reviews()[0]['evidence']
        self.assertFalse(evidence['recoverable'])
        self.assertFalse(evidence['releasable'])
        with self.assertRaises(ValueError):
            self.session.review_action('1','recover',True)
        self.assertEqual(load_queue(self.queue).iloc[0]['Status'],'REVISAO')

    def test_unsigned_and_incomplete_journal_never_become_eligible(self):
        records={'1':{},'2':{'phase':'unknown'},'3':{'phase':'rejected','error':{'category':'limit_quota','api_code':'rate_limit_exceeded'}}}
        self.set_statuses(['PENDENTE','PENDENTE','ERRO'])
        plan=self.selection('pending_errors',records)
        self.assertNotIn('1',plan['selected_ids'])
        self.assertNotIn('2',plan['selected_ids'])
        self.assertNotIn('3',plan['selected_ids'])

    def test_simulation_linked_history_reopens_preserves_old_and_blocks_api(self):
        before=self.queue.read_bytes()
        with patch('socket.socket',side_effect=AssertionError('Network forbidden')):
            with self.assertRaisesRegex(ValueError,'bloqueada'):
                self.session.start(real=True)
            self.session.start()
            self.session.worker.join(10)
            original=deepcopy(self.session.run_log)
            self.session.set_selection('resume',original['run_id'])
            self.session.start()
            self.session.worker.join(10)
        self.assertEqual(self.queue.read_bytes(),before)
        self.assertFalse(self.session.allow_api)
        runs=self.session.history()['runs']
        self.assertEqual(len(runs),2)
        self.assertEqual(self.session.run_log['parent_run_id'],original['run_id'])
        self.assertEqual(next(r for r in runs if r['run_id']==original['run_id']),original)
        fresh=StudioSession(self.root/'fresh.json',env={})
        fresh.choose_queue(self.queue)
        self.assertEqual(len(fresh.history()['runs']),2)

    def test_invalid_history_isolated_and_empty_state(self):
        archive=Path(str(self.queue)+'.studio-runs');archive.mkdir()
        (archive/'broken.json').write_text('{','utf-8')
        (archive/'bad-shape.json').write_text('{"events":[1]}','utf-8')
        (archive/'old.json').write_text('{"version":1,"run_id":"old","events":[]}','utf-8')
        history=self.session.history()
        self.assertEqual(len(history['warnings']),2)
        self.assertEqual(len(history['runs']),1)
        self.assertNotIn('selected_ids',history['runs'][0])
        empty=StudioSession(self.root/'empty.json',env={})
        self.assertEqual(empty.history(),{'runs':[],'warnings':[]})

    def test_malformed_optional_history_fields_do_not_break_interface(self):
        archive=Path(str(self.queue)+'.studio-runs');archive.mkdir()
        malformed=[{'run_id':[]}, {'events':[{'outputs':1}]}, {'events':[{'diagnostic':[]}]},
                   {'settings':{'image':1}}, {'selected_ids':'1,2'}]
        for i, value in enumerate(malformed):
            (archive/f'bad-{i}.json').write_text(json.dumps(value),'utf-8')
        history=self.session.history()
        self.assertEqual(history['runs'],[])
        self.assertEqual(len(history['warnings']),len(malformed))
        self.session.choose_queue(self.queue)
        self.assertIsNone(self.session.run_log)

    def test_release_invalidated_after_reference_changes(self):
        self.set_statuses(['REVISAO'])
        job=self.prepare_record()
        self.session.review_action('1','release',True)
        job.referencias[0].write_bytes(b'changed reference')
        plan=self.selection('review',Journal(self.queue).records)
        self.assertNotIn('1',plan['selected_ids'])
        self.assertEqual(plan['decisions'][0]['category'],'review')

    def test_resume_keeps_authorized_reviews_and_pending_inside_original_scope(self):
        self.set_statuses(['REVISAO'])
        self.prepare_record('1');self.session.review_action('1','release',True)
        self.set_statuses(['REVISAO'])
        plan=self.selection('resume',Journal(self.queue).records,['1','2'])
        self.assertEqual(plan['selected_ids'],['1','2'])
        self.assertTrue(all(d['category']=='outside' for d in plan['decisions'][2:]))
