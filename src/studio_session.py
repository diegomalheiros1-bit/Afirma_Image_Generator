"""Local Studio state and worker; no browser storage or credentials in settings."""
from copy import deepcopy
from dataclasses import asdict
from datetime import datetime, timezone
import json
import math
import os
from pathlib import Path
import re
import tempfile
import threading
import time
import uuid

from PIL import Image
from src.app_info import app_info
from src.credential_store import CredentialStore
from src.excel_reader import load_queue
from src.execution_state import queue_lock, PersistenceError, digest
from src.image_settings import ImageSettings
from src.job_values import cell_text, parse_quantity


def validate_photos(paths):
    resolved = tuple(Path(p).resolve(strict=True) for p in paths)
    if not 1 <= len(resolved) <= 16:
        raise ValueError("Selecione entre 1 e 16 fotos de referência.")
    for path in resolved:
        if not path.is_file() or path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            raise ValueError(f"Referência não suportada: {path.name}")
        if path.stat().st_size >= 50 * 1024 * 1024:
            raise ValueError(f"Referência deve ter menos de 50 MB: {path.name}")
        try:
            with Image.open(path) as image:
                if image.format not in {"PNG", "JPEG", "WEBP"}:
                    raise ValueError()
                image.verify()
        except Exception:
            raise ValueError(f"Conteúdo de imagem inválido: {path.name}") from None
    return resolved


def validate_folders(folders, default):
    result = {}
    for entry in folders:
        alias = entry.get("alias", "").strip()
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,39}", alias):
            raise ValueError("Alias inválido: comece com letra; use até 40 letras, números, _ ou -.")
        if alias.casefold() in {a.casefold() for a in result}:
            raise ValueError(f"Alias duplicado: {alias}")
        raw = entry.get("path", "").strip()
        if not raw or not Path(raw).is_dir():
            raise ValueError(f"Pasta inexistente: {alias} ({raw})")
        result[alias] = Path(raw).resolve()
    if default and default not in result:
        raise ValueError("Pasta padrão não cadastrada.")
    return result


def save_folder_sheet(queue, folders, default):
    """Explicit user operation only; atomic write preserves other workbook sheets."""
    from openpyxl import load_workbook
    from src.excel_reader import _save_workbook_atomic
    roots = validate_folders(folders, default)
    with queue_lock(queue):
        book = load_workbook(queue)
        try:
            if "Pastas_Referencias" in book:
                del book["Pastas_Referencias"]
            sheet = book.create_sheet("Pastas_Referencias")
            sheet.append(["Alias", "Caminho"])
            for alias, path in roots.items():
                sheet.append([alias, str(path)])
            if default:
                config = (book["Configuracao"] if "Configuracao" in book
                          else book.create_sheet("Configuracao"))
                if config.max_row == 1 and config['A1'].value is None:
                    config['A1'], config['B1'] = "Parametro", "Valor"
                for row in config.iter_rows(min_row=2):
                    if row[0].value == "Pasta_Referencias":
                        row[1].value = str(roots[default])
                        break
                else:
                    config.append(["Pasta_Referencias", str(roots[default])])
            _save_workbook_atomic(book, queue)
        finally:
            book.close()


class RunControl:
    def __init__(self, callback, simulation_delay=0):
        self.condition = threading.Condition()
        self.paused = False
        self.at_pause = False
        self.stopped = False
        self.callback = callback
        self.delay = simulation_delay

    def event(self, kind, **values):
        self.callback(kind, values)

    def checkpoint(self):
        if not self._wait_pause():
            return False
        if self.delay:
            time.sleep(self.delay)
        return self._wait_pause()

    def _wait_pause(self):
        with self.condition:
            if not self.paused or self.stopped:
                self.at_pause = False
                return not self.stopped
            self.at_pause = True
        # Never take the session lock while holding the condition: commands
        # acquire these locks in the opposite order.
        self.event('paused')
        with self.condition:
            while self.paused and not self.stopped:
                self.condition.wait(.25)
            self.at_pause = False
            return not self.stopped

    def command(self, action):
        with self.condition:
            if action == "pause":
                self.paused = True
            elif action == "resume":
                if not self.paused or not self.at_pause:
                    raise ValueError("Aguarde a pausa efetiva antes de retomar.")
                self.paused = False
            elif action == "stop":
                self.stopped = True
            else:
                raise ValueError("Comando de execução inválido.")
            self.condition.notify_all()


class StudioSession:
    def __init__(self, settings_path, *, allow_api=False, env=None):
        self.lock = threading.RLock()
        self.settings_path = Path(settings_path)
        self.allow_api = allow_api
        self.app_mode = "browser"
        self.env = env  # Defaults are extracted locally; secrets never enter state/preferences.
        self.credential_file = Path(__file__).resolve().parents[1] / ".env"
        self.credential_store = CredentialStore(self.settings_path)
        self._session_api_key = None
        self._credential_revision = 0
        self.connection_worker = None
        self._api_connection = dict(status="unverified", message="Verifique a conexão para confirmar sua chave.", checked_at=None)
        self.queue = None
        self.queue_summary = None
        self.photos = []
        self.folders = []
        self.default_folder = ""
        self.mode = "spreadsheet"
        self.preferences = dict(image=asdict(ImageSettings()), output_dir="", timeout=120, retries=2,
                                max_jobs=1, max_images=1, safe_mode=True)
        if env is None:
            from dotenv import dotenv_values
            defaults = {**dotenv_values(self.credential_file), **os.environ}
        else:
            defaults = env
        self._external_api_key_available = bool(defaults.get("OPENAI_API_KEY"))
        self.preferences.update(image=asdict(ImageSettings.from_env(defaults)),
                                timeout=float(defaults.get("OPENAI_TIMEOUT_SECONDS", 120)),
                                retries=int(defaults.get("OPENAI_RETRIES", 2)),
                                max_jobs=int(defaults.get("MAX_JOBS_PER_RUN", 1)),
                                max_images=int(defaults.get("MAX_IMAGES_PER_RUN", 1)))
        from main import flag
        self.preferences["safe_mode"] = flag(defaults.get("FIRST_RUN_SAFE_MODE", "SIM"))
        self.checked_preferences(self.preferences)
        if self.settings_path.exists():
            self._load_preferences(json.loads(self.settings_path.read_text("utf-8")))
        self.active = False
        self.control = None
        self.worker = None
        self.result = {}
        self.items = []
        self.total = 0
        self.status = "Pronto"
        self.error = ""
        self.run_log = None
        self._run_clock = None
        self._item_clocks = {}
        self.selection_mode = 'pending'
        self.resume_run_id = None
        self.explicit_ids = []

    def history(self):
        if not self.queue:
            return {'runs': [], 'warnings': []}
        paths = list(Path(str(self.queue) + '.studio-runs').glob('*.json'))
        paths.append(Path(str(self.queue) + '.studio-run.json'))
        records, warnings = {}, []
        for path in paths:
            if not path.exists():
                continue
            try:
                value = json.loads(path.read_text('utf-8'))
                if not isinstance(value, dict) or not isinstance(value.get('events', []), list):
                    raise ValueError()
                if any(not isinstance(e, dict) for e in value.get('events', [])):
                    raise ValueError()
                if value.get('run_id') is not None and not isinstance(value['run_id'], str):
                    raise ValueError()
                for event in value.get('events', []):
                    if 'outputs' in event and (not isinstance(event['outputs'], list) or
                            any(not isinstance(path, str) for path in event['outputs'])):
                        raise ValueError()
                    for field in ('diagnostic', 'usage'):
                        if field in event and not isinstance(event[field], dict):
                            raise ValueError()
                for field in ('settings', 'counts', 'result'):
                    if field in value and not isinstance(value[field], dict):
                        raise ValueError()
                for field in ('selected_ids', 'original_selected_ids'):
                    if field in value and (not isinstance(value[field], list) or
                            any(not isinstance(key, str) or not key for key in value[field])):
                        raise ValueError()
                if 'image' in value.get('settings', {}) and not isinstance(value['settings']['image'], dict):
                    raise ValueError()
                records[value.get('run_id') or path.name] = value
            except (ValueError, OSError):
                warnings.append('Histórico indisponível ou inválido: ' + path.name)
        return {'runs': sorted(records.values(), key=lambda r: str(r.get('started_at', '')), reverse=True), 'warnings': warnings}

    def set_selection(self, mode, run_id=None, ids=None):
        with self.lock:
            from src.run_selection import MODES
            from src.job_values import canonical_id
            self.idle()
            if mode not in MODES:
                raise ValueError('Seleção inválida.')
            self.selection_mode, self.resume_run_id = mode, run_id
            self.explicit_ids = [canonical_id(v) for v in (ids or [])]

    def selection(self, config, prefs):
        from src.run_selection import plan_selection
        from src.execution_state import Journal
        original = None
        if self.selection_mode == 'resume':
            previous = next((r for r in self.history()['runs'] if r.get('run_id') == self.resume_run_id), None)
            if self.resume_run_id and previous is None:
                raise ValueError('Execução de origem indisponível. Selecione outro histórico ou informe os IDs explicitamente.')
            previous = previous or self.run_log or {}
            original = previous.get('original_selected_ids') or previous.get('selected_ids') or self.explicit_ids
        frame = load_queue(self.queue, reference_mode=self.mode)
        return plan_selection(frame, Journal(self.queue).records, config, prefs['max_jobs'], prefs['max_images'], self.selection_mode, original)

    def reviews(self):
        from src.execution_state import Journal
        from src.run_selection import review_evidence
        from src.job_values import canonical_id
        if not self.queue:
            return []
        config, _ = self.snapshot()
        records = Journal(self.queue).records
        result = []
        for _, row in load_queue(self.queue, reference_mode=self.mode).iterrows():
            key = canonical_id(row['ID'])
            record = records.get(key, {})
            if cell_text(row['Status']).upper() not in {'REVISAO', 'PROCESSANDO'} and record.get('phase') not in {'uncertain', 'in_flight', 'retry_authorized'}:
                continue
            from src.generation_errors import sanitize_diagnostic
            event = next((e for run in self.history()['runs'] for e in reversed(run.get('events', []))
                          if e.get('id') == key and e.get('type') == 'item_end'), {})
            result.append(dict(id=key, released=record.get('phase') == 'retry_authorized',
                diagnostic=sanitize_diagnostic(record.get('error', {})),
                calls=record.get('api_attempts'), at=event.get('at'), elapsed_seconds=event.get('elapsed_seconds'),
                evidence=review_evidence(row, record, config)))
        return result

    def review_action(self, item_id, action, consent=False):
        with self.lock:
            from src.execution_state import Journal
            from src.run_selection import review_evidence
            from src.job_values import canonical_id
            from src.excel_reader import save_queue
            from src.generation_errors import sanitize_diagnostic
            self.idle()
            if action not in {'recover', 'release'} or consent is not True:
                raise ValueError('Confirme explicitamente a ação de revisão.')
            config, _ = self.snapshot()
            with queue_lock(self.queue):
                frame = load_queue(self.queue, reference_mode=self.mode)
                index = next((i for i, r in frame.iterrows() if canonical_id(r['ID']) == item_id), None)
                if index is None:
                    raise ValueError('Item não encontrado.')
                journal = Journal(self.queue)
                record = journal.records.get(item_id, {})
                if cell_text(frame.loc[index, 'Status']).upper() not in {'REVISAO', 'PROCESSANDO'}:
                    raise ValueError('Este item não está em revisão.')
                evidence = review_evidence(frame.loc[index], record, config)
                if action == 'recover':
                    if not evidence['recoverable']:
                        raise ValueError(evidence['reason'])
                    frame.at[index, 'Status'] = 'CONCLUIDO'
                    frame.at[index, 'Observacao'] = 'Recuperado por ação explícita; hashes e identidade comprovados.'
                    journal.put(item_id, review_actions=list(record.get('review_actions', [])) +
                                [dict(at=self._stamp(), action='recover', verified_identity=True, verified_hashes=True)])
                    save_queue(frame, self.queue)
                else:
                    if not evidence['releasable']:
                        raise ValueError(evidence['reason'])
                    audit = list(record.get('review_actions', []))
                    audit.append(dict(at=self._stamp(), action='release', prior_phase=record.get('phase'),
                                      diagnostic=sanitize_diagnostic(record.get('error', {})), duplicate_charge_acknowledged=True,
                                      prior_attempt={name: deepcopy(record.get(name)) for name in
                                          ('execution', 'execution_fingerprint', 'row_fingerprint', 'row_fingerprint_version',
                                           'outputs', 'hashes', 'api_attempts', 'task_attempts')}))
                    journal.put(item_id, phase='retry_authorized', review_actions=audit)
                    frame.at[index, 'Status'] = 'PENDENTE'
                    frame.at[index, 'Observacao'] = 'Nova tentativa liberada explicitamente; possível cobrança adicional reconhecida.'
                    save_queue(frame, self.queue)
            self._refresh_queue_summary()

    @staticmethod
    def _stamp():
        return datetime.now(timezone.utc).isoformat(timespec="seconds")

    def _save_run_log(self):
        """Keep a credential-free, atomic record beside the selected workbook."""
        if not self.run_log or not self.queue:
            return
        path = Path(str(self.queue) + ".studio-run.json")
        paths = [path]
        if self.run_log.get("run_id"):
            archive = Path(str(self.queue) + ".studio-runs")
            archive.mkdir(exist_ok=True)
            paths.insert(0, archive / (self.run_log["run_id"] + ".json"))
        for destination in paths:
            temporary = None
            try:
                with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=destination.parent,
                                                 prefix=".studio-run-", delete=False) as handle:
                    temporary = Path(handle.name)
                    json.dump(self.run_log, handle, ensure_ascii=False, indent=2)
                    handle.flush()
                    os.fsync(handle.fileno())
                os.replace(temporary, destination)
            finally:
                if temporary is not None:
                    temporary.unlink(missing_ok=True)

    def idle(self):
        if self.active:
            raise ValueError("Execução ativa: pause/retome ou encerre antes de alterar a campanha.")

    def api_key_status(self):
        if self._session_api_key:
            return "session"
        if self.credential_store.exists():
            return "saved"
        return "external" if self._external_api_key_available else "missing"

    def set_api_key(self, key, *, persist=False):
        with self.lock:
            self.idle()
            if self._api_connection["status"] == "checking":
                raise ValueError("Aguarde a verificação terminar antes de alterar a chave.")
            if type(key) is not str or type(persist) is not bool:
                raise ValueError("Informe uma chave válida.")
            key = key.strip()
            if not 20 <= len(key) <= 512 or any(character.isspace() or ord(character) < 32 for character in key):
                raise ValueError("Informe uma chave de API válida, sem espaços.")
            if persist:
                try:
                    self.credential_store.save(key)
                except OSError:
                    raise ValueError("Não foi possível salvar a chave protegida. A chave anterior foi mantida.") from None
                self._session_api_key = None
            else:
                self._session_api_key = key
            self._reset_api_connection()

    def clear_api_key(self, source=None):
        with self.lock:
            self.idle()
            if self._api_connection["status"] == "checking":
                raise ValueError("Aguarde a verificação terminar antes de alterar a chave.")
            source = source or self.api_key_status()
            if source == "saved":
                try:
                    self.credential_store.clear()
                except OSError:
                    raise ValueError("Não foi possível remover a chave salva. Tente novamente.") from None
            elif source == "session":
                self._session_api_key = None
            else:
                raise ValueError("A configuração externa é gerenciada fora do Studio.")
            self._reset_api_connection()

    def resolve_api_key(self, external=None):
        return self._session_api_key or self.credential_store.read() or external

    def _reset_api_connection(self):
        self._credential_revision += 1
        self._api_connection = dict(status="unverified", message="Verifique a conexão para confirmar sua chave.", checked_at=None)

    def verify_api_key(self):
        """Start one bounded probe without blocking UI polling or generating images."""
        with self.lock:
            self.idle()
            if self._api_connection["status"] == "checking":
                return deepcopy(self._api_connection)
            try:
                if self.env is None:
                    from dotenv import dotenv_values
                    external = {**dotenv_values(self.credential_file), **os.environ}
                else:
                    external = self.env
                self._external_api_key_available = bool(external.get("OPENAI_API_KEY"))
                key = self.resolve_api_key(external.get("OPENAI_API_KEY"))
            except OSError:
                self._api_connection = dict(status="disconnected", message="Não foi possível ler a chave protegida. Informe sua chave novamente.", checked_at=None)
                return deepcopy(self._api_connection)
            if not key:
                self._api_connection = dict(status="disconnected", message="Informe uma chave para conectar sua conta.", checked_at=None)
                return deepcopy(self._api_connection)
            revision = self._credential_revision
            self._api_connection = dict(status="checking", message="Confirmando a autenticação com a OpenAI…", checked_at=None)

            def work():
                from src.api_connection import check_api_connection
                result = check_api_connection(key)
                with self.lock:
                    # A late response for a replaced/removed key must never turn
                    # the new credential green or overwrite its current check.
                    if revision == self._credential_revision:
                        self._api_connection = {**result, "checked_at": self._stamp()}

            self.connection_worker = threading.Thread(target=work, name="afirma-key-check", daemon=True)
            self.connection_worker.start()
            return deepcopy(self._api_connection)

    @staticmethod
    def checked_preferences(values):
        allowed = {"image", "output_dir", "timeout", "retries", "max_jobs", "max_images", "safe_mode"}
        if set(values) - allowed:
            raise ValueError("Configuração desconhecida; credenciais não são aceitas pelo Studio.")
        ImageSettings(**values["image"])
        if not isinstance(values.get("output_dir"), str):
            raise ValueError("Pasta de saída inválida.")
        if values["output_dir"] and not Path(values["output_dir"]).is_absolute():
            raise ValueError("Escolha uma pasta de saída absoluta.")
        if type(values["timeout"]) not in (int, float) or not math.isfinite(values["timeout"]) or values["timeout"] <= 0:
            raise ValueError("Tempo limite deve ser positivo.")
        for key in ("max_jobs", "max_images"):
            if type(values[key]) is not int or values[key] < 1:
                raise ValueError("Limites devem ser inteiros positivos.")
        if type(values["retries"]) is not int or not 0 <= values["retries"] <= 5:
            raise ValueError("Novas tentativas: escolha de 0 a 5.")
        if type(values["safe_mode"]) is not bool:
            raise ValueError("Modo seguro inválido.")
        return deepcopy(values)

    def _load_preferences(self, data):
        if set(data) - {"version", "preferences"} or data.get("version") != 1:
            raise ValueError("Arquivo de preferências inválido ou versão desconhecida.")
        self.preferences = self.checked_preferences(data["preferences"])

    def set_preferences(self, values, persist=False):
        with self.lock:
            self.idle()
            values = self.checked_preferences(values)
            if persist:
                self.settings_path.parent.mkdir(parents=True, exist_ok=True)
                temporary = None
                try:
                    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=self.settings_path.parent,
                                                     delete=False) as handle:
                        temporary = Path(handle.name)
                        json.dump({"version": 1, "preferences": values}, handle, indent=2)
                        handle.flush()
                        os.fsync(handle.fileno())
                    os.replace(temporary, self.settings_path)
                finally:
                    if temporary:
                        temporary.unlink(missing_ok=True)
            self.preferences = values
            if self.queue:
                self._refresh_queue_summary()

    def _refresh_queue_summary(self, config=None):
        if not self.queue:
            self.queue_summary = None
            return
        frame = load_queue(self.queue, reference_mode=self.mode)
        prompts = frame.apply(
            lambda row: bool(cell_text(row.get("Prompt_Padrao")) or
                             cell_text(row.get("Prompt_Variacao"))), axis=1)
        pending = frame["Status"].map(lambda value: isinstance(value, str) and value.strip().upper() == "PENDENTE")
        pending_prompts = prompts & pending
        image_count = 0
        invalid_quantities = 0
        execution_items = 0
        execution_images = 0
        # The Studio form is the explicit source of execution limits. The
        # workbook limit remains a CLI default and must not silently cap it.
        item_limit = self.preferences["max_jobs"]
        for _, row in frame.loc[pending_prompts].iterrows():
            try:
                quantity = parse_quantity(row.get("Quantidade"))
                image_count += quantity
            except ValueError:
                invalid_quantities += 1
                continue
            if execution_items >= item_limit or execution_images + quantity > self.preferences["max_images"]:
                continue
            execution_items += 1
            execution_images += quantity
        completed = frame["Status"].map(
            lambda value: isinstance(value, str) and value.strip().upper() in {"CONCLUIDO", "SIMULADO"})
        completed_real = frame["Status"].map(
            lambda value: isinstance(value, str) and value.strip().upper() == "CONCLUIDO")
        completed_images = 0
        for _, row in frame.loc[prompts & completed_real].iterrows():
            try:
                completed_images += parse_quantity(row.get("Quantidade"))
            except ValueError:
                continue
        review = frame["Status"].map(
            lambda value: isinstance(value, str) and value.strip().upper() in {"REVISAO", "ERRO"})
        self.queue_summary = {
            "items_with_prompt": int(prompts.sum()),
            "pending_items": int(pending_prompts.sum()),
            "estimated_images": image_count,
            "execution_items": execution_items,
            "execution_images": execution_images,
            "item_limit": item_limit,
            "image_limit": self.preferences["max_images"],
            "invalid_quantities": invalid_quantities,
            "completed_items": int((prompts & completed).sum()),
            "completed_images": completed_images,
            "review_items": int((prompts & review).sum()),
        }

    def choose_queue(self, path):
        with self.lock:
            self.idle()
            path = Path(path).resolve(strict=True)
            previous = (self.queue, self.queue_summary)
            self.queue = path
            try:
                self._refresh_queue_summary()
            except Exception:
                self.queue, self.queue_summary = previous
                raise
            self.items, self.result, self.error = [], {}, ""
            self.selection_mode, self.resume_run_id, self.explicit_ids = "pending", None, []
            report = Path(str(path) + ".studio-run.json")
            try:
                self.run_log = next(iter(self.history()['runs']), None)
            except (ValueError, OSError):
                self.run_log = None

    def set_mode(self, mode):
        with self.lock:
            self.idle()
            if mode not in {"spreadsheet", "direct"}:
                raise ValueError("Modo inválido.")
            previous = self.mode
            self.mode = mode
            try:
                if self.queue:
                    self._refresh_queue_summary()
            except Exception:
                self.mode = previous
                raise
    def add_photos(self, paths):
        with self.lock:
            self.idle()
            additions = validate_photos(paths)
            proposed = list(self.photos)
            for p in additions:
                if p not in proposed:  # deduplicate exact resolved paths, never basenames.
                    proposed.append(p)
            validate_photos(proposed)
            self.photos = proposed

    def remove_photo(self, index):
        with self.lock:
            self.idle()
            if type(index) is not int or not 0 <= index < len(self.photos):
                raise ValueError("Foto não encontrada na seleção.")
            self.photos.pop(index)

    def set_folders(self, folders, default):
        with self.lock:
            self.idle()
            validate_folders(folders, default)
            if self.default_folder and not default and self.folders:
                raise ValueError("Escolha outra pasta padrão antes de remover a pasta padrão atual. Arquivos sem alias não serão redirecionados automaticamente.")
            self.folders, self.default_folder = deepcopy(folders), default

    def folder_usage(self, alias):
        with self.lock:
            self.idle()
            if not self.queue:
                return []
            df = load_queue(self.queue, reference_mode="direct")
            used = []
            for _, row in df.iterrows():
                for ref in cell_text(row["Arquivo_Referencia"]).split(";"):
                    if (("::" in ref and ref.split("::", 1)[0].strip() == alias) or
                            alias == self.default_folder and ref.strip() and "::" not in ref):
                        used.append(str(row["ID"]))
                        break
            return used

    def save_folders(self):
        with self.lock:
            self.idle()
            if not self.queue:
                raise ValueError("Selecione uma planilha antes de salvar as pastas.")
            save_folder_sheet(self.queue, self.folders, self.default_folder)

    def snapshot(self):
        if not self.queue:
            raise ValueError("Selecione a planilha da campanha.")
        prefs = self.checked_preferences(self.preferences)
        if not prefs["output_dir"]:
            raise ValueError("Escolha a pasta de saída no Studio antes de conferir ou gerar.")
        # The menu is the only configuration source in Studio. CLI keeps its
        # workbook configuration contract for existing command-line workflows.
        config = dict(studio=True, reference_mode=self.mode, limit=prefs["max_jobs"],
                      dry_run=True, reference_dirs={}, reference_dir=None,
                      result_dir=Path(prefs["output_dir"]).resolve(),
                      image_settings=ImageSettings(**prefs["image"]))
        if self.mode == "spreadsheet":
            roots = validate_folders(self.folders, self.default_folder)
            config["reference_dirs"] = roots
            config["reference_dir"] = roots[self.default_folder] if self.default_folder else None
            if config["reference_dir"] is None and not roots:
                raise ValueError("Escolha uma pasta padrão para usar referências pela planilha.")
        else:
            config["direct_references"] = validate_photos(self.photos)
        return config, prefs

    def validate_campaign(self):
        from main import make_job
        from src.image_generator import OpenAIImageGenerator
        from src.job_values import canonical_id
        config, prefs = self.snapshot()
        df = load_queue(self.queue, reference_mode=self.mode)
        ids = df["ID"].map(canonical_id)
        if ids.eq("").any() or ids.duplicated().any():
            raise ValueError("Cada linha deve ter ID preenchido, único e estável.")
        from src.execution_state import Journal
        records = Journal(self.queue).records
        selection = self.selection(config, prefs)
        errors, plan, outputs = [], [], set()
        for _, row in df.iterrows():
            if not any(d['id'] == canonical_id(row['ID']) and d['category'] in {'selected', 'limit'} for d in selection['decisions']):
                continue
            try:
                job = make_job(row, config, True)
                OpenAIImageGenerator.validate(job)
                if any(p in outputs for p in job.saidas):
                    raise ValueError("Nome de saída repetido em outra linha da campanha.")
                outputs.update(job.saidas)
                if any(p.exists() or p.is_symlink() for p in job.saidas):
                    # Recovery can prove completed outputs; only block unregistered conflicts here.
                    from src.execution_state import Journal
                    if str(job.id) not in Journal(self.queue).records:
                        raise ValueError("Saída já existe sem histórico confirmado; escolha um nome novo.")
                plan.append({"id": str(job.id), "prompt": job.prompt,
                             "references": [str(p) for p in job.referencias],
                             "outputs": [str(p) for p in job.saidas], "quantity": job.quantidade})
            except (ValueError, OSError) as exc:
                errors.append(f"ID {row['ID']}: {exc}")
        if errors:
            raise ValueError("\n".join(errors))
        item_limit = min(prefs["max_jobs"], config["limit"])
        selected, images = 0, 0
        for job in plan:
            if selected >= item_limit or images + job["quantity"] > prefs["max_images"]:
                continue
            selected += 1
            images += job["quantity"]
        return {"jobs": plan, "settings": asdict(config["image_settings"]),
                "experimental": config["image_settings"].experimental,
                "selection": {k: v for k, v in selection.items() if k != "indices"},
                "execution_preview": dict(items=len(selection['selected_ids']), images=selection['images'],
                    output_dir=str(config["result_dir"]),
                    safe_mode_compatible=not prefs["safe_mode"] or (
                        prefs["max_jobs"] <= 1 and prefs["max_images"] <= 1 and config["limit"] <= 1))}

    def _event(self, kind, values):
        with self.lock:
            if kind == 'paused':
                if self.active and self.control.paused and self.control.at_pause and not self.control.stopped:
                    self.status = 'Pausado'
                    self.run_log['status'] = self.status
            elif kind == "plan":
                self.total = values["total"]
                self.run_log["planned_items"] = values["total"]
                self.run_log["planned_images"] = values["images"]
            elif kind == "item_start":
                self._item_clocks[values["id"]] = time.monotonic()
                self.run_log["events"].append({"at": self._stamp(), "type": "item_start", "id": values["id"]})
            elif kind == "item":
                self.items.append(values)
                started = self._item_clocks.pop(values["id"], None)
                event = {"at": self._stamp(), "type": "item_end", "id": values["id"],
                    "status": values["status"], "message": values["message"],
                    "elapsed_seconds": round(time.monotonic() - started, 3) if started is not None else None}
                for key in ("outputs", "api_attempts", "usage", "request_id", "diagnostic"):
                    if values.get(key) is not None:
                        event[key] = values[key]
                self.run_log["events"].append(event)
            if self.run_log:
                try:
                    self._save_run_log()
                except OSError:
                    self.run_log["log_error"] = "Falha ao persistir o log; preserve a planilha e confira o disco."

    def start(self, real=False):
        with self.lock:
            self.idle()
            if real and not self.allow_api:
                raise ValueError("API real bloqueada nesta sessão. A implementação deve ser validada antes da autorização.")
            plan = self.validate_campaign()
            config, prefs = self.snapshot()
            if not plan['execution_preview']['items']:
                raise ValueError('Nenhum item elegível. Confira os motivos na prévia e os itens em revisão.')
            config['selection_mode'] = self.selection_mode
            config['original_ids'] = plan['selection']['selected_ids']
            config['selection_ids'] = plan['selection']['selected_ids']
            config["dry_run"] = not real
            config["expected_reference_hashes"] = {
                path: digest(Path(path)) for job in plan["jobs"] for path in job["references"]}
            queue = self.queue
            self.active = True
            self.error, self.result, self.items, self.total = "", {}, [], plan['execution_preview']['items']
            self._run_clock = time.monotonic()
            self._item_clocks = {}
            previous = next((r for r in self.history()['runs'] if r.get('run_id') == self.resume_run_id), self.run_log or {}) if self.selection_mode == 'resume' else {}
            self.run_log = {"version": 2, 'selected_ids': plan['selection']['selected_ids'],
                            'original_selected_ids': previous.get('original_selected_ids') or previous.get('selected_ids') or plan['selection']['selected_ids'],
                            'parent_run_id': previous.get('run_id'), 'selection_mode': self.selection_mode,
                            'selection_decisions': plan['selection']['decisions'], "run_id": uuid.uuid4().hex, "started_at": self._stamp(), "finished_at": None,
                            "elapsed_seconds": None, "mode": "real" if real else "simulation",
                            "output_dir": plan['execution_preview']['output_dir'],
                            "status": "Gerando" if real else "Simulando", "planned_items": plan['execution_preview']['items'],
                            "planned_images": plan['execution_preview']['images'],
                            "settings": {"image": deepcopy(prefs["image"]), "max_jobs": prefs["max_jobs"],
                                         "max_images": prefs["max_images"], "timeout": prefs["timeout"],
                                         "retries": prefs["retries"]}, "events": [], "result": {}}
            try:
                self._save_run_log()
            except OSError:
                self.run_log["log_error"] = "Falha ao persistir o log; preserve a planilha e confira o disco."
            self.status = "Gerando" if real else "Simulando"
            self.control = RunControl(self._event, simulation_delay=0 if real else 0.25)
            control = self.control
            environment = dict(self.env) if self.env is not None else None
            session_api_key = self._session_api_key

            def work():
                from main import main
                final_status = 'Erro'
                try:
                    if environment is None and real:
                        from dotenv import dotenv_values
                        effective_env = {**dotenv_values(self.credential_file), **os.environ}
                    else:
                        effective_env = dict(environment or {})
                    if real:
                        selected_key = session_api_key or self.credential_store.read() or effective_env.get("OPENAI_API_KEY")
                        if selected_key:
                            effective_env["OPENAI_API_KEY"] = selected_key
                    effective_env.update(IMAGE_PROVIDER="openai" if real else "simulation", DRY_RUN="NAO" if real else "SIM",
                                         FIRST_RUN_SAFE_MODE="SIM" if prefs["safe_mode"] else "NAO",
                                         MAX_JOBS_PER_RUN=str(prefs["max_jobs"]), MAX_IMAGES_PER_RUN=str(prefs["max_images"]),
                                         OPENAI_TIMEOUT_SECONDS=str(prefs["timeout"]), OPENAI_RETRIES=str(prefs["retries"]))
                    result = main(queue, env=effective_env, overrides=config, control=control)
                    with self.lock:
                        self.result = result
                    final_status = result.get('stop_reason') or ('Finalizado com erros' if result['errors'] else 'Finalizado')
                except Exception as exc:
                    with self.lock:
                        self.error = str(exc) if isinstance(exc, (ValueError, OSError, PersistenceError)) else "Falha local; confira o log sem apagar o histórico."
                finally:
                    with self.lock:
                        try:
                            self._refresh_queue_summary()
                        except Exception:
                            self.queue_summary = None
                            if not self.error:
                                self.error = "Execução encerrada; não foi possível atualizar o resumo da planilha. Selecione-a novamente para conferir os estados."
                        self.run_log["finished_at"] = self._stamp()
                        self.run_log["elapsed_seconds"] = round(time.monotonic() - self._run_clock, 3)
                        self.run_log["status"] = final_status
                        self.run_log["result"] = dict(self.result)
                        started = sum(event['type'] == 'item_start' for event in self.run_log['events'])
                        self.run_log['counts'] = dict(planned=self.run_log['planned_items'], started=started,
                            completed=self.result.get('success', 0),
                            review=sum(event.get('status') == 'REVISAO' for event in self.run_log['events']),
                            waiting=max(0, self.run_log['planned_items'] - started))
                        self.run_log['counts']['errors'] = sum(e.get('status') == 'ERRO' for e in self.run_log['events'])
                        self.run_log['selection_counts'] = {category: sum(d['category'] == category for d in self.run_log['selection_decisions'])
                            for category in ('outside', 'limit', 'completed', 'error', 'review')}
                        durations = [e['elapsed_seconds'] for e in self.run_log['events'] if e.get('status') in {'CONCLUIDO', 'SIMULADO'} and isinstance(e.get('elapsed_seconds'), (int, float))]
                        self.run_log['mean_completed_seconds'] = sum(durations)/len(durations) if durations else None
                        if self.error:
                            self.run_log["error"] = self.error
                        try:
                            self._save_run_log()
                        except OSError:
                            self.run_log["log_error"] = "Falha ao persistir o log; preserve a planilha e confira o disco."
                        # Publish completion only after the last persistence
                        # attempt, under the same lock used by state().
                        self.status = final_status
                        self.active = False

            self.worker = threading.Thread(target=work, name="afirma-campaign", daemon=False)
            self.worker.start()
            return plan

    def open_output_folder(self):
        """Open only the configured local output directory, including while running."""
        with self.lock:
            if self.active and self.run_log and self.run_log.get('output_dir'):
                path = Path(self.run_log['output_dir'])
            else:
                config, _ = self.snapshot()
                path = Path(config['result_dir'])
            path = path.resolve()
            if not path.is_dir():
                raise ValueError('A pasta de resultados ainda não existe. Ela será criada ao gerar imagens.')
            if os.name != 'nt':
                raise ValueError('A abertura de pasta está disponível no Windows.')
            os.startfile(str(path), 'explore')

    def command(self, action):
        with self.lock:
            if not self.active or not self.control:
                raise ValueError("Nenhuma execução ativa.")
            if self.control.stopped and action != 'stop':
                raise ValueError('Interrupção já solicitada; aguarde o item atual e o salvamento do histórico.')
            if action == 'pause' and self.control.paused:
                return
            self.control.command(action)
            self.status = {"pause": "Pausa solicitada — após o item atual", "resume": "Executando",
                           "stop": "Interrupção solicitada — aguardando o item atual"}[action]
            if self.run_log:
                self.run_log['status'] = self.status
                if action == 'stop':
                    self.run_log['stop_requested_at'] = self._stamp()
                try:
                    self._save_run_log()
                except OSError:
                    self.run_log['log_error'] = 'Falha ao persistir o log; preserve a planilha e confira o disco.'

    def state(self):
        with self.lock:
            return dict(queue=str(self.queue) if self.queue else "", queue_summary=deepcopy(self.queue_summary), mode=self.mode,
                        photos=[{"name": p.name, "path": str(p)} for p in self.photos],
                        folders=deepcopy(self.folders), default_folder=self.default_folder,
                        preferences=deepcopy(self.preferences), active=self.active,
                        paused=bool(self.active and self.control and self.control.paused and self.control.at_pause and not self.control.stopped), status=self.status,
                        pause_requested=bool(self.active and self.control and self.control.paused and not self.control.stopped),
                        stopping=bool(self.active and self.control and self.control.stopped),
                        total=self.total, items=deepcopy(self.items), result=dict(self.result), error=self.error,
                        run_log=deepcopy(self.run_log), history=self.history(),
                        selection_mode=self.selection_mode, resume_run_id=self.resume_run_id,
                        current_item_seconds=(round(time.monotonic()-min(self._item_clocks.values()), 3) if self.active and self._item_clocks else None),
                        allow_api=self.allow_api, app_mode=self.app_mode, app_info=app_info(), api_key_status=self.api_key_status(),
                        api_connection=deepcopy(self._api_connection))
