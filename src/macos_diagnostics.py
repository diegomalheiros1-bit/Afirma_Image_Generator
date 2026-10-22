"""Isolated WKWebView smoke test; no credentials or paid generation."""
import json
from pathlib import Path
import socket
import time


def prepare_diagnostics(root):
    from openpyxl import Workbook, load_workbook
    from PIL import Image
    from src.studio_session import StudioSession
    import openai

    root = Path(root)
    queue = root / 'diagnostic.xlsx'
    if queue.exists():
        raise ValueError('Use uma pasta nova para diagnóstico.')
    for method_name in ('connect', 'connect_ex'):
        original = getattr(socket.socket, method_name)

        def guard(instance, address, method=original):
            if not isinstance(address, tuple) or address[0] not in {'127.0.0.1', 'localhost', '::1'}:
                raise AssertionError('External network forbidden')
            return method(instance, address)
        setattr(socket.socket, method_name, guard)

    def forbidden(*args, **kwargs):
        raise AssertionError('API forbidden in diagnostics')
    openai.OpenAI = forbidden
    book = Workbook()
    sheet = book.active
    sheet.title = 'Fila_Geracao'
    sheet.append(['ID', 'Prompt_Padrao', 'Prompt_Variacao', 'Arquivo_Referencia',
                  'Nome_Saida', 'Status', 'Tentativas', 'Observacao'])
    sheet.append([1, 'Produto em fundo branco.', '', 'reference.png', 'teste.png', 'PENDENTE', 0, ''])
    book.save(queue)
    book.close()
    Image.new('RGB', (32, 32), 'purple').save(root / 'reference.png')
    session = StudioSession(root / 'settings.json', allow_api=False, env={})

    class NoCredentials:
        def exists(self):
            return False

        def read(self):
            return None
    session.credential_store = NoCredentials()
    session.set_folders([{'alias': 'Produtos', 'path': str(root)}], 'Produtos')
    session.set_preferences({**session.preferences, 'output_dir': str(root / 'images'),
                             'max_jobs': 1, 'max_images': 1, 'safe_mode': False})
    session.choose_queue(queue)

    def ready(app):
        report = {'passed': False, 'checks': [], 'engine': app.window.gui.renderer}
        def evaluate(expression):
            # evaluate_js wraps code in eval(), blocked by the app CSP in
            # WKWebView. Native run_js evaluates directly, preserving the CSP.
            return json.loads(app.window.run_js(f'JSON.stringify({expression})'))
        try:
            assert report['engine'] == 'wkwebview'
            deadline = time.monotonic() + 30
            while not evaluate("typeof state !== 'undefined' && state !== null"):
                if time.monotonic() > deadline:
                    raise AssertionError('UI did not initialize')
                time.sleep(.2)
            assert evaluate("q('error').getBoundingClientRect().height === 0")
            assert evaluate("state.app_mode === 'desktop' && !state.allow_api")
            report['checks'].append('wkwebview-loaded-empty-error-hidden-api-disabled')
            for width, height in [(1280, 800), (720, 700), (520, 650)]:
                app.window.resize(width, height)
                time.sleep(.4)
                assert evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
            report['checks'].append('responsive-layout-three-sizes')
            session.start(real=False)
            session.worker.join(60)
            assert not session.active
            result = session.run_log.get('result') or {}
            assert result.get('success') == 1, result
            report['checks'].append('isolated-excel-reference-simulation')
            with queue.open('rb') as stream:
                check = load_workbook(stream, read_only=True)
                assert check['Fila_Geracao'].max_row == 2
                check.close()
            report['passed'] = True
        except Exception as exc:
            report['failure'] = f'{type(exc).__name__}: {exc}'
        finally:
            (root / 'diagnostics.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
            app.window.destroy()
    return session, ready
