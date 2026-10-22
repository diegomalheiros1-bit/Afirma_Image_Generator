"""Opt-in native diagnostic using its own ten-item workbook and fake keys."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import time


def prepare_diagnostics(root):
    from openpyxl import Workbook
    from PIL import Image
    from src.studio_session import StudioSession

    root = Path(root)
    queue, reference = root / "diagnostic-campaign.xlsx", root / "diagnostic-reference.png"
    if queue.exists():
        raise ValueError("Use uma nova pasta para o diagnóstico; arquivos existentes são preservados.")
    Image.new("RGB", (32, 32), "purple").save(reference)
    book = Workbook()
    sheet = book.active
    sheet.title = "Fila_Geracao"
    sheet.append(["ID", "Prompt_Padrao", "Prompt_Variacao", "Arquivo_Referencia", "Nome_Saida", "Status", "Tentativas", "Observacao"])
    for i in range(1, 11):
        sheet.append([str(i), "Fotografia de produto.", "Teste isolado.", reference.name,
                      f"produto.{i}", "PENDENTE", 0, ""])
    config = book.create_sheet("Configuracao")
    for row in [("Parametro", "Valor"), ("Pasta_Referencias", str(root)),
                ("Pasta_Resultados", str(root / "images")), ("Limite_Por_Execucao", 1), ("Dry_Run", "SIM")]:
        config.append(row)
    book.save(queue)
    book.close()
    import openai
    import socket
    original_connect = socket.socket.connect
    original_connect_ex = socket.socket.connect_ex
    def guard(method):
        def connect(instance, address):
            if not isinstance(address, tuple) or address[0] not in {'127.0.0.1', 'localhost', '::1'}:
                raise AssertionError('External network forbidden in diagnostics')
            return method(instance, address)
        return connect
    socket.socket.connect = guard(original_connect)
    socket.socket.connect_ex = guard(original_connect_ex)
    def forbidden_client(*args, **kwargs):
        raise AssertionError('Network API client forbidden in diagnostics')
    openai.OpenAI = forbidden_client
    session = StudioSession(root / "diagnostic-settings.json", allow_api=False, env={})
    values = deepcopy(session.preferences)
    values.update(output_dir=str(root / "images"), max_jobs=10, max_images=10, safe_mode=False)
    session.set_preferences(values)
    session.set_folders([{'alias': 'Produtos', 'path': str(root)}], 'Produtos')

    def ready(app):
        window = app.window
        report = {"passed": False, "engine": window.gui.renderer, "layouts": [], "checks": []}
        before = hashlib.sha256(queue.read_bytes()).hexdigest()

        def wait(script, timeout=15):
            deadline = time.monotonic() + timeout
            while time.monotonic() < deadline:
                result = window.evaluate_js(script)
                if result:
                    return result
                time.sleep(.1)
            raise AssertionError("A interface não alcançou o estado esperado.")

        def click(selector):
            wait(f"(() => {{const e=document.querySelector({json.dumps(selector)});const r=e?.getBoundingClientRect();return r?.width>0&&r?.height>0&&!e.disabled;}})()")
            window.evaluate_js(f"document.querySelector({json.dumps(selector)}).click()")

        def capture(name):
            # Capture the actual WebView2 renderer, including CSS and native DPI.
            # DOM evaluation can finish before WebView2 paints its next frame.
            time.sleep(.2)
            from System import Action
            from System.IO import FileStream, FileMode, FileAccess
            from Microsoft.Web.WebView2.Core import CoreWebView2CapturePreviewImageFormat
            stream = FileStream(str(root / f"{name}.png"), FileMode.Create, FileAccess.Write)
            value = {}
            try:
                def start():
                    value["task"] = window.native.browser.webview.CoreWebView2.CapturePreviewAsync(
                        CoreWebView2CapturePreviewImageFormat.Png, stream)
                window.native.Invoke(Action(start))
                value["task"].GetAwaiter().GetResult()
            finally:
                stream.Close()

        def layout(name):
            value = window.evaluate_js("""(() => {
              const bad=[...document.querySelectorAll('button,input,select,a')].filter(e=>{
                const r=e.getBoundingClientRect();return r.width&&r.height&&(r.left < -1 || r.right > innerWidth+1);
              }).map(e=>e.id||e.tagName);
              return {width:innerWidth,height:innerHeight,scrollWidth:document.documentElement.scrollWidth,
                scrollY:window.scrollY, columns:getComputedStyle(document.querySelector('.grid')).gridTemplateColumns,
                clippedControls:bad};
            })()""")
            assert value["scrollWidth"] <= value["width"] + 1, value
            assert not value["clippedControls"], value
            report["layouts"].append({"name": name, **value})
            capture(name)

        try:
            assert window.gui.renderer == "edgechromium"
            wait("typeof state !== 'undefined' && state !== null")
            from src.app_info import APP_VERSION
            report['app_version'] = window.evaluate_js('state.app_info.version')
            assert report['app_version'] == APP_VERSION
            assert window.evaluate_js('state.app_mode') == 'desktop'
            assert window.evaluate_js("q('error').getBoundingClientRect().height===0")
            window.evaluate_js("report(new Error('Aviso fictício de validação.'))")
            assert window.evaluate_js("q('error').getBoundingClientRect().height>0 && q('error').textContent.includes('Aviso fictício')")
            window.evaluate_js("q('error').textContent=''")
            assert window.evaluate_js("q('error').getBoundingClientRect().height===0")
            assert window.evaluate_js("[...document.querySelectorAll('.error,.notice')].filter(e=>!e.textContent.trim()).every(e=>getComputedStyle(e).display==='none')")
            report['checks'].append('empty-messages-take-no-space-show-on-error-and-hide-after-clear')
            assert window.evaluate_js("q('history-data').textContent.includes('Nenhuma execução')")
            assert window.evaluate_js("!q('persistent-history').open")
            click('#history-heading')
            assert window.evaluate_js("q('persistent-history').open")
            window.evaluate_js("q('persistent-history').scrollIntoView({block:'center'})")
            layout('history-empty-expanded')
            click('#history-heading')
            layout('empty-campaign-history')
            app.pick_files = lambda kind: {"queue": [str(queue)], "photos": [str(reference)], "folder": [str(root / 'images')]}[kind]
            click("#choose-queue")
            wait("state.queue_summary && state.queue_summary.items_with_prompt === 10")
            assert window.evaluate_js("state.queue_summary.execution_images") == 10
            report["checks"].append("workbook-selection")
            window.evaluate_js("document.querySelector('#reference-mode').value='direct';document.querySelector('#reference-mode').dispatchEvent(new Event('change',{bubbles:true}))")
            wait("state.mode === 'direct'")
            click("#add-photos")
            wait("state.photos.length === 1")
            click("#choose-output")
            wait("state.preferences.output_dir.length > 0")
            report["checks"].append("reference-and-output-selection")
            layout("home-default")
            default_width = report["layouts"][-1]["width"]
            window.maximize()
            wait(f"innerWidth > {default_width}")
            assert str(window.native.WindowState) == "Maximized"
            layout("home-maximized")
            window.restore()
            report["checks"].append("native-maximize-and-restore")
            for width, height in [(1024, 768), (850, 640), (700, 600), (520, 500)]:
                window.resize(width, height)
                wait(f"innerWidth <= {width} && innerWidth >= {width - 40}")
                layout(f"home-{width}")
            click("#settings-button")
            preferences_before_tabs = deepcopy(session.preferences)
            click("#settings-tab-about")
            wait("document.querySelector('#settings-tab-about').getAttribute('aria-selected') === 'true'")
            assert window.evaluate_js("document.querySelector('#about-version').textContent") == 'Versão ' + session.state()['app_info']['version']
            assert window.evaluate_js("document.querySelector('#about-model').textContent") == session.preferences['image']['model']
            assert window.evaluate_js("[...document.querySelectorAll('.about-links a')].every(a=>a.target==='_blank'&&a.rel.includes('noopener')&&a.rel.includes('noreferrer'))")
            window.evaluate_js("document.querySelector('.settings-tabs').scrollIntoView({block:'start'})")
            layout("about-small-light")
            window.evaluate_js("document.querySelector('#settings-tab-about').dispatchEvent(new KeyboardEvent('keydown',{key:'ArrowLeft',bubbles:true}))")
            wait("document.activeElement.id === 'settings-tab-generation'")
            window.evaluate_js("document.querySelector('#timeout').value='123';document.querySelector('#timeout').dispatchEvent(new Event('input',{bubbles:true}))")
            window.evaluate_js("document.querySelector('#model').value='gpt-image-2.5-flare';document.querySelector('#model').dispatchEvent(new Event('input',{bubbles:true}))")
            click("#settings-tab-about")
            assert window.evaluate_js("document.querySelector('#about-model').textContent") == 'gpt-image-2.5-flare'
            time.sleep(1)
            click("#settings-tab-generation")
            assert window.evaluate_js("document.querySelector('#timeout').value") == '123'
            assert window.evaluate_js("document.querySelector('#model').value") == 'gpt-image-2.5-flare'
            assert session.preferences == preferences_before_tabs
            window.evaluate_js("document.querySelector('#model').value='gpt-image-2.5-sunburst';document.querySelector('#model').dispatchEvent(new Event('input',{bubbles:true}))")
            window.evaluate_js("document.querySelector('#timeout').value='120';document.querySelector('#timeout').dispatchEvent(new Event('input',{bubbles:true}))")
            report["checks"].append("about-tab-keyboard-links-and-unsaved-preferences")
            click("#settings details > summary")
            wait("document.querySelector('#settings details').open")
            layout("settings-small")
            # Test error placement with a deliberately incomplete, local-only key.
            window.evaluate_js("document.querySelector('#api-key').value='curta';document.querySelector('#api-key').dispatchEvent(new Event('input',{bubbles:true}))")
            click("#api-key-session")
            wait("document.querySelector('#api-key-error').textContent.length > 0")
            assert session.api_key_status() == "missing"
            report["checks"].append("incomplete-key-error")
            window.evaluate_js("document.querySelector('.credential-card').scrollIntoView({block:'start'})")
            layout("api-key-small-error")
            window.evaluate_js("document.querySelector('#api-key').value='';document.querySelector('#api-key').dispatchEvent(new Event('input',{bubbles:true}))")
            click("#theme-toggle")
            window.evaluate_js("document.querySelector('.credential-card').scrollIntoView({block:'start'})")
            layout("settings-small-dark")
            window.resize(1200, 800)
            wait("innerWidth > 1100")
            layout("settings-wide-dark")
            window.evaluate_js("document.querySelector('.credential-card').scrollIntoView({block:'center'})")
            layout("api-key-wide-dark")
            # Credential state flow, entirely fictional and local. DPAPI is real,
            # provider probes are fakes and the external network is blocked.
            from unittest.mock import patch
            import threading
            fake_key = 'fictional-diagnostic-credential-only'
            def type_key():
                window.evaluate_js("q('api-key').value='fictional-diagnostic-credential-only';q('api-key').dispatchEvent(new Event('input',{bubbles:true}))")
            def card_layout(name):
                window.evaluate_js("document.querySelector('.credential-card').scrollIntoView({block:'start'})")
                layout(name)
            with patch('src.api_connection.check_api_connection') as probe:
                type_key()
                click('#api-key-save')
                wait("state.api_key_status==='saved' && q('credential-form').classList.contains('hidden')")
                assert session.resolve_api_key() == fake_key
                assert fake_key not in json.dumps(session.state())
                assert fake_key.encode() not in session.credential_store.path.read_bytes()
                reopened = StudioSession(session.settings_path, env={})
                assert reopened.resolve_api_key() == fake_key
                assert reopened.state()['api_connection']['status'] == 'unverified'
                assert window.evaluate_js("q('api-key').value==='' && q('api-key').type==='password'")
                card_layout('credential-saved-wide-dark')
                click('#api-key-change')
                wait("document.activeElement.id==='api-key'")
                type_key()
                click('#api-key-toggle')
                assert window.evaluate_js("q('api-key').type==='text'")
                click('#api-key-cancel')
                assert window.evaluate_js("document.activeElement.id==='api-key-change' && q('api-key').value==='' && q('api-key').type==='password'")
                click('#api-key-change')
                type_key()
                with patch.object(session.credential_store, 'save', side_effect=OSError(fake_key)):
                    click('#api-key-save')
                    wait("q('api-key-error').textContent.includes('anterior')")
                assert session.resolve_api_key() == fake_key
                assert window.evaluate_js("q('api-key').value===''")
                assert fake_key not in window.evaluate_js('document.body.innerText')
                card_layout('credential-storage-failure-dark')
                click('#api-key-cancel')
                probe.assert_not_called()
            finish = threading.Event()
            def fake_probe(key):
                assert key == fake_key
                finish.wait(10)
                return dict(status='connected', message='Autenticação fictícia confirmada.')
            with patch('src.api_connection.check_api_connection', side_effect=fake_probe) as probe:
                click('#api-key-check')
                wait("state.api_connection.status==='checking'")
                assert window.evaluate_js("q('api-key-check').disabled && q('api-key-change').disabled && q('api-key-clear').disabled")
                session.verify_api_key()
                assert probe.call_count == 1
                card_layout('credential-checking-dark')
                finish.set()
                wait("state.api_connection.status==='connected'")
            click('#theme-toggle')
            card_layout('credential-verified-wide-light')
            window.resize(520, 600)
            wait('innerWidth < 520')
            card_layout('credential-saved-small-light')
            click('#api-key-change')
            card_layout('credential-replacement-small-light')
            type_key()
            click('#api-key-session')
            wait("state.api_key_status==='session' && state.api_connection.status==='unverified'")
            card_layout('credential-session-small-light')
            # Test actual scoped HTTP removal without an interactive native prompt.
            window.evaluate_js("window.originalCredentialConfirm=window.confirm;window.confirm=()=>true")
            click('#api-key-clear')
            wait("state.api_key_status==='saved'")
            assert 'salva' in window.evaluate_js("q('credential-notice').textContent")
            session._external_api_key_available = True
            session.env = {'OPENAI_API_KEY': 'fictional-external-diagnostic-only'}
            click('#api-key-clear')
            wait("state.api_key_status==='external'")
            assert window.evaluate_js("q('api-key-clear').classList.contains('hidden')")
            card_layout('credential-external-small-light')
            click('#api-key-change')
            type_key()
            click('#back')
            assert window.evaluate_js("q('api-key').value===''")
            click('#settings-button')
            assert window.evaluate_js("q('credential-form').classList.contains('hidden')")
            session._external_api_key_available = False
            session.env = {}
            window.evaluate_js("window.confirm=window.originalCredentialConfirm;refresh()")
            wait("state.api_key_status==='missing'")
            assert window.evaluate_js("!q('credential-form').classList.contains('hidden')")
            card_layout('credential-missing-small-light')
            assert session.connection_worker is None or not session.connection_worker.is_alive()
            report['checks'].append('credential-lifecycle-no-network-dpapi-reopen-focus-removal-fallback')
            click('#theme-toggle')
            window.resize(1200, 800)
            wait('innerWidth > 1100')
            from src.desktop_app import on_ui_thread
            for factor in (1.25, 1.5):
                on_ui_thread(window, lambda: setattr(window.native.browser.webview, 'ZoomFactor', factor))
                wait(f"innerWidth < {1200 / factor}")
                window.evaluate_js("document.querySelector('.credential-card').scrollIntoView({block:'start'})")
                layout(f"settings-zoom-{int(factor * 100)}")
            on_ui_thread(window, lambda: setattr(window.native.browser.webview, 'ZoomFactor', 1.0))
            wait("innerWidth > 1100")
            report["checks"].append("webview-zoom-125-and-150-percent")
            click("#settings-tab-about")
            window.evaluate_js("document.querySelector('.settings-tabs').scrollIntoView({block:'start'})")
            layout("about-wide-dark")
            window.resize(520, 500)
            wait("innerWidth < 520")
            window.evaluate_js("document.querySelector('.settings-tabs').scrollIntoView({block:'start'})")
            layout("about-small-dark")
            window.evaluate_js("document.querySelector('.about-resources').scrollIntoView({block:'start'})")
            layout("about-links-small-dark")
            window.resize(1200, 800)
            wait("innerWidth > 1100")
            click("#settings-tab-generation")
            click("#back")
            click('#settings-button')
            window.evaluate_js("q('quality').value='max';q('timeout').value=120;q('quality').dispatchEvent(new Event('input',{bubbles:true}));q('timeout').dispatchEvent(new Event('input',{bubbles:true}))")
            wait("!q('timeout-guidance').classList.contains('hidden')")
            assert window.evaluate_js("q('timeout').value") == '120'
            window.evaluate_js("q('timeout-guidance').scrollIntoView({block:'center'})")
            layout('timeout-guidance-wide-dark')
            window.evaluate_js("q('timeout').value=600;q('timeout').dispatchEvent(new Event('input',{bubbles:true}))")
            wait("q('timeout-guidance').classList.contains('hidden')")
            window.evaluate_js("q('timeout').value=120;q('timeout').dispatchEvent(new Event('input',{bubbles:true}))")
            click('#apply')
            wait("state.preferences.image.quality==='max' && state.preferences.timeout===120")
            assert session.preferences['timeout'] == 120
            click('#back')
            report['checks'].append('max-timeout-guidance-nonblocking-no-silent-preference-changes')
            assert window.evaluate_js("!q('validate') && !q('plan-dialog')")
            click("#selection-preview")
            wait("q('selection-preview-data').textContent.includes('10 itens selecionados')")
            assert window.evaluate_js("q('queue-summary').textContent.includes('10 itens nesta execução')")
            window.evaluate_js("window.diagnosticLimitButton=document.querySelector('.summary-limits');diagnosticLimitButton.focus()")
            time.sleep(.9)
            assert window.evaluate_js("document.activeElement===diagnosticLimitButton && document.querySelector('.summary-limits')===diagnosticLimitButton")
            assert window.evaluate_js("q('selection-preview-data').textContent.includes('10 itens selecionados')")
            assert all(Path(job['outputs'][0]).suffix == '.png'
                       for job in session.validate_campaign()['jobs'])
            click(".summary-limits")
            wait("document.activeElement === q('max-jobs')")
            assert window.evaluate_js("!q('settings').classList.contains('hidden')")
            click("#back")
            report["checks"].append("automatic-preflight-selection-count-and-limits-shortcut")
            # UI-only paid-flow fixture: the backend keeps allow_api=False and
            # every real start is intercepted before reaching the server.
            window.evaluate_js("""
              window.diagnosticOriginalApi=api;window.diagnosticPaidCalls=0;window.diagnosticKeyAvailable=false;
              api=async function(path,body){
                if(path==='start'&&body?.real){window.diagnosticPaidCalls++;await new Promise(resolve=>setTimeout(resolve,100));return {};}
                const result=await window.diagnosticOriginalApi(path,body);
                if(path==='state'){result.allow_api=true;result.api_key_status=window.diagnosticKeyAvailable?'session':'missing';}
                return result;
              };refresh();
            """)
            wait("!document.querySelector('#real-start').classList.contains('hidden')")
            window.resize(520, 500)
            wait("innerWidth < 520")
            for button in ('selection-preview', 'start', 'real-start'):
                window.evaluate_js(f"document.querySelector('#{button}').scrollIntoView({{block:'center'}});document.querySelector('#{button}').focus()")
                wait("!document.querySelector('#action-tooltip').classList.contains('hidden')")
                assert window.evaluate_js("document.querySelector('#action-tooltip').textContent.length") > 40
            layout("action-tooltip-small-dark")
            click("#real-start")
            wait("document.querySelector('#generation-dialog').open")
            assert window.evaluate_js("document.querySelector('#generation-confirm').disabled")
            layout("generation-missing-key-dark")
            click("#generation-cancel")
            assert window.evaluate_js("window.diagnosticPaidCalls") == 0
            window.evaluate_js("window.diagnosticKeyAvailable=true;refresh()")
            wait("state.api_key_status === 'session'")
            window.resize(1200, 800)
            wait("innerWidth > 1100")
            click("#real-start")
            wait("document.querySelector('#generation-dialog').open")
            assert window.evaluate_js("document.querySelector('#generation-images').textContent") == 'Até 10'
            assert window.evaluate_js("document.querySelector('#generation-items').textContent") == 'Até 10'
            assert window.evaluate_js("!q('generation-timeout-guidance').classList.contains('hidden') && !q('generation-confirm').disabled")
            layout("generation-wide-dark")
            window.evaluate_js("""window.diagnosticRegenerationPlan=JSON.parse(JSON.stringify(generationPlan));
                diagnosticRegenerationPlan.selection.decisions[0].regeneration=true;
                renderGenerationDialog(diagnosticRegenerationPlan);""")
            assert window.evaluate_js("!q('generation-regeneration').classList.contains('hidden') && q('generation-regeneration').textContent.includes('IDs 1:') && !q('generation-confirm').disabled")
            layout('generation-regeneration-wide-dark')
            window.resize(520, 500)
            wait("innerWidth < 520")
            layout('generation-regeneration-small-dark')
            capture('generation-regeneration-small-dark')
            window.evaluate_js("renderGenerationDialog(generationPlan)")
            assert window.evaluate_js("q('generation-regeneration').classList.contains('hidden')")
            report['checks'].append('pending-regeneration-notice-wide-small-no-api')
            layout("generation-small-dark")
            window.evaluate_js("document.querySelector('#generation-dialog').dispatchEvent(new Event('cancel',{cancelable:true}))")
            wait("!document.querySelector('#generation-dialog').open")
            assert window.evaluate_js("window.diagnosticPaidCalls") == 0
            click("#theme-toggle")
            click("#real-start")
            wait("document.querySelector('#generation-dialog').open")
            layout("generation-small-light")
            # A changed preview requires a fresh confirmation, with zero starts.
            session.preferences['max_images'] = 2
            click("#generation-confirm")
            wait("document.querySelector('#generation-error').textContent.includes('mudou')")
            assert window.evaluate_js("window.diagnosticPaidCalls") == 0
            assert window.evaluate_js("document.querySelector('#generation-images').textContent") == 'Até 2'
            window.evaluate_js("document.querySelector('#generation-confirm').click();document.querySelector('#generation-confirm').click()")
            wait("!document.querySelector('#generation-dialog').open")
            assert window.evaluate_js("window.diagnosticPaidCalls") == 1
            assert not session.active and not session.allow_api
            assert hashlib.sha256(queue.read_bytes()).hexdigest() == before
            session.preferences['max_images'] = 10
            window.evaluate_js("api=window.diagnosticOriginalApi;refresh()")
            wait("state.allow_api === false && state.api_key_status === 'missing'")
            report["checks"].append("paid-dialog-cancel-escape-missing-key-stale-plan-and-single-submit-no-api")
            window.resize(1200, 800)
            wait("innerWidth > 1100")
            click("#start")
            wait("state.result.success === 10")
            assert hashlib.sha256(queue.read_bytes()).hexdigest() == before
            assert not Path(str(queue) + ".state.json").exists()
            report["checks"].append("ten-item-studio-limit-overrides-sheet-one-workbook-unchanged")
            assert window.evaluate_js("q('success').textContent === '10' && q('errors').textContent === '0' && q('remaining-count').textContent === '0'")
            assert window.evaluate_js("q('success-label').textContent === 'Simulados'")
            assert window.evaluate_js("q('bar').parentElement.getAttribute('aria-valuenow') === '100'")
            layout("simulation-finished")
            window.evaluate_js("q('history-panel').open=true")
            window.evaluate_js("document.querySelector('.run-history').scrollIntoView({block:'start'})")
            layout("log-wide-dark")
            window.resize(520, 500)
            wait("innerWidth < 520")
            window.evaluate_js("document.querySelector('.log-table-wrap').scrollIntoView({block:'center'})")
            layout("log-small-dark")
            assert window.evaluate_js("document.querySelector('.log-table-wrap').scrollWidth > document.querySelector('.log-table-wrap').clientWidth")
            report["checks"].append("log-scroll-contained-in-table")
            # Keep the real gate off for the entire native diagnostic.
            assert session.allow_api is False
            try:
                session.start(real=True)
                raise AssertionError('Real gate must reject starts')
            except ValueError:
                pass
            assert window.evaluate_js("[0,59,60,61,119,120,3600,59.5,null].map(formatDuration).join('|')") == '0 s|59 s|1 min|1 min 1 s|1 min 59 s|2 min|1 h|1 min|—'
            assert window.evaluate_js("[59.4,119.5,3920,undefined,NaN,Infinity,-1].map(formatDuration).join('|')") == '59 s|2 min|1 h 5 min 20 s|—|—|—|—'
            report['checks'].append('shared-duration-boundaries-real-gate-blocked')
            # Disclosure regressions: only fixture files and existing state requests.
            assert window.evaluate_js("!q('persistent-history').open")
            window.evaluate_js("window.historyNode=q('history-data').firstChild;window.historyPanel=q('persistent-history');window.historyApi=api;window.historyCalls=[];api=async(...args)=>{historyCalls.push(args[0]);return historyApi(...args)}")
            for theme in ('light', 'dark'):
                window.evaluate_js(f"document.documentElement.dataset.theme='{theme}';themeLabel();q('persistent-history').open=false;q('persistent-history').scrollIntoView({{block:'center'}})")
                layout(f'history-collapsed-{theme}')
                click('#history-heading')
                assert window.evaluate_js("q('persistent-history').open")
                layout(f'history-expanded-{theme}')
            time.sleep(1.2)
            assert window.evaluate_js("q('persistent-history').open && historyPanel===q('persistent-history') && historyNode===q('history-data').firstChild")
            assert window.evaluate_js("historyCalls.every(x=>x==='state')")
            window.evaluate_js("api=historyApi;document.documentElement.dataset.theme='light';themeLabel()")
            window.maximize()
            time.sleep(.3)
            for theme in ('light', 'dark'):
                for expanded in (False, True):
                    window.evaluate_js(f"document.documentElement.dataset.theme='{theme}';themeLabel();q('persistent-history').open={str(expanded).lower()};q('persistent-history').scrollIntoView({{block:'center'}})")
                    layout(f'history-maximized-{theme}-{"expanded" if expanded else "collapsed"}')
            window.restore()
            window.resize(520,800)
            wait('innerWidth<600')
            window.evaluate_js("document.documentElement.dataset.theme='light';themeLabel();q('history-heading').scrollIntoView({block:'center'})")
            # Real keyboard events delivered to this diagnostic WebView2 only.
            from System import Action
            def key(key, code, vk):
                for event_type in ('keyDown', 'keyUp'):
                    params=json.dumps(dict(type=event_type, key=key, code=code, windowsVirtualKeyCode=vk,
                        text=('\r' if key=='Enter' else ' ' if key==' ' else '') if event_type=='keyDown' else ''))
                    task={}
                    def send():
                        task['value']=window.native.browser.webview.CoreWebView2.CallDevToolsProtocolMethodAsync('Input.dispatchKeyEvent',params)
                    window.native.Invoke(Action(send))
                    task['value'].GetAwaiter().GetResult()
            window.evaluate_js("q('history-heading').focus()")
            key('Enter','Enter',13)
            wait("!q('persistent-history').open")
            key(' ','Space',32)
            wait("q('persistent-history').open")
            key('Tab','Tab',9)
            assert window.evaluate_js("document.activeElement===q('history-choice')")
            # Same-data polling retains both the DOM and user disclosure state.
            window.evaluate_js("q('persistent-history').open=false")
            time.sleep(1.2)
            assert window.evaluate_js("!q('persistent-history').open")
            # Use a second temporary workbook through the actual session API.
            from shutil import copyfile
            second_queue=root/'second-campaign.xlsx'
            copyfile(queue,second_queue)
            session.choose_queue(second_queue)
            wait("state.queue.endsWith('second-campaign.xlsx')")
            assert window.evaluate_js("!q('persistent-history').open && q('history-data').textContent.includes('Nenhuma')")
            click('#history-heading')
            session.choose_queue(queue)
            wait("state.queue.endsWith('diagnostic-campaign.xlsx')")
            assert window.evaluate_js("!q('persistent-history').open")
            click('#history-heading')
            # Multiple-run fixture rendered without creating or altering stored runs.
            window.evaluate_js("window.savedHistory=structuredClone(state.history);const extra=structuredClone(state.history.runs[0]);extra.run_id='fixture-older-run';extra.status='Fixture anterior';state.history.runs.push(extra);renderPersistentHistory();q('history-choice').value='fixture-older-run';q('history-choice').dispatchEvent(new Event('change'));document.querySelector('#history-data details').open=true;window.selectedHistoryNode=q('history-data').firstChild;renderPersistentHistory()")
            assert window.evaluate_js("q('history-choice').value==='fixture-older-run' && q('history-data').textContent.includes('Fixture anterior') && document.querySelector('#history-data details').open && selectedHistoryNode===q('history-data').firstChild && q('persistent-history').open")
            window.evaluate_js("state.history.runs[0].status='Fixture atualizada';renderPersistentHistory()")
            assert window.evaluate_js("selectedHistoryNode===q('history-data').firstChild && document.querySelector('#history-data details').open")
            window.evaluate_js("state.history.runs[1].status='Fixture anterior atualizada';renderPersistentHistory()")
            assert window.evaluate_js("document.querySelector('#history-data details').open && q('persistent-history').open && q('history-choice').value==='fixture-older-run'")
            layout('history-multiple-runs')
            window.evaluate_js("state.history=savedHistory;renderPersistentHistory();q('history-heading').focus()")
            report['checks'].append('history-native-disclosure-keyboard-polling-dom-preserved-queue-reset-empty-multiple-runs-no-extra-api')
            window.evaluate_js("q('persistent-history').scrollIntoView({block:'start'})")
            wait("q('history-data').textContent.includes('Simulado') || q('history-data').textContent.includes('Itens planejados')")
            layout('persistent-history-small')
            click('#selection-preview')
            wait("q('selection-preview-data').textContent.includes('10 itens')")
            window.evaluate_js("q('selection-preview-data').querySelector('details').open=true;q('execution-selection').scrollIntoView({block:'start'})")
            layout('selection-preview-small')
            window.resize(1200, 800)
            wait('innerWidth>1100')
            window.evaluate_js("q('persistent-history').scrollIntoView({block:'start'})")
            layout('persistent-history-wide')
            reopened = StudioSession(root / 'reopened-settings.json', allow_api=False, env={})
            reopened.choose_queue(queue)
            assert len(reopened.history()['runs']) == 1
            window.evaluate_js("[...q('history-data').querySelectorAll('button')].find(b=>b.textContent==='Retomar lote').click()")
            wait("state.selection_mode==='resume'")
            click('#selection-preview')
            wait("q('selection-preview-data').textContent.includes('10 itens')")
            assert session.selection_mode == 'resume'
            assert session.allow_api is False
            assert window.evaluate_js("[...q('execution-selection').options].map(o=>o.textContent).join('|')")=='Gerar pendentes|Repetir erros|Repetir itens em revisão'
            assert window.evaluate_js("q('execution-choice').classList.contains('hidden') && !q('resume-options').classList.contains('hidden')")
            layout('resume-from-history-small')
            window.evaluate_js("document.activeElement?.blur();[...q('history-data').querySelectorAll('button')].find(b=>b.textContent==='Retomar lote').scrollIntoView({block:'center'})")
            # Allow scrolling/layout to finish before a fresh keyboard focus.
            # Focusing an already focused button after click does not emit focusin.
            time.sleep(.25)
            window.evaluate_js("[...q('history-data').querySelectorAll('button')].find(b=>b.textContent==='Retomar lote').focus({preventScroll:true})")
            wait("!q('action-tooltip').classList.contains('hidden') && q('action-tooltip').textContent.includes('execução original')")
            layout('history-resume-tooltip')
            click('#resume-exit')
            wait("state.selection_mode==='pending' && !q('execution-choice').classList.contains('hidden')")
            window.evaluate_js("[...q('history-data').querySelectorAll('button')].find(b=>b.textContent==='Retomar lote').click()")
            wait("state.selection_mode==='resume'")
            report['checks'].append('three-actions-resume-only-from-history')
            report['checks'].append('history-reopened-original-selection-preview-api-remains-blocked')
            click('#load-reviews')
            wait("q('review-data').textContent.includes('Nenhum item')")
            window.evaluate_js("q('review-section').scrollIntoView({block:'start'})")
            layout('review-empty-wide')
            # Fictional interrupted batch: six completed, one uncertain, three not started.
            from src.execution_state import Journal, execution_identity, row_fingerprint
            from src.excel_reader import load_queue
            from main import make_job
            book = __import__('openpyxl').load_workbook(queue)
            for i in range(2, 8):
                book['Fila_Geracao'][f'F{i}'] = 'CONCLUIDO'
            book['Fila_Geracao']['F8'] = 'REVISAO'
            book.save(queue)
            book.close()
            config, prefs = session.snapshot()
            row = load_queue(queue, reference_mode='direct').iloc[6]
            job = make_job(row, config, True)
            identity, signature = execution_identity(job, row)
            signature_row = dict(row)
            signature_row['Arquivo_Referencia'] = ''
            Journal(queue).put('7', phase='uncertain', outputs=[str(p) for p in job.saidas],
                execution=identity, execution_fingerprint=signature, row_fingerprint=row_fingerprint(signature_row),
                row_fingerprint_version=2, api_attempts=1,
                error=dict(category='connection', exception_type='APIConnectionError', uncertain=True, request_id='req_fictional'))
            click('#selection-preview')
            wait("q('selection-preview-data').textContent.includes('3 itens')")
            click('#load-reviews')
            wait("q('review-data').textContent.includes('Item #7')")
            window.evaluate_js("q('review-section').scrollIntoView({block:'start'})")
            layout('review-uncertain-wide')
            window.resize(520,500)
            wait('innerWidth<520')
            window.evaluate_js("q('review-section').scrollIntoView({block:'start'})")
            layout('review-uncertain-small')
            window.evaluate_js("[...q('review-data').querySelectorAll('button')].find(b=>b.textContent.includes('Liberar')).click()")
            wait("[...document.querySelectorAll('dialog[open]')].some(d=>d.textContent.includes('outra cobrança'))")
            assert window.evaluate_js("q('action-tooltip').textContent.includes('Fecha a confirmação')")
            layout('review-confirmation-small')
            window.evaluate_js("[...document.querySelector('dialog[open]').querySelectorAll('button')].find(b=>b.textContent==='Cancelar').click()")
            assert load_queue(queue, reference_mode='direct').iloc[6]['Status'] == 'REVISAO'
            assert session.allow_api is False
            report['checks'].append('fictional-six-completed-one-review-three-resume-review-consent-cancel')
            window.evaluate_js("[...q('review-data').querySelectorAll('button')].find(b=>b.textContent.includes('Liberar')).click()")
            wait("[...document.querySelectorAll('dialog[open]')].some(d=>d.textContent.includes('outra cobrança'))")
            window.evaluate_js("[...document.querySelector('dialog[open]').querySelectorAll('button')].find(b=>b.textContent==='Confirmar').click()")
            wait("state.selection_mode==='review' && q('review-feedback').textContent.includes('Item #7 liberado') && q('selection-preview-data').textContent.includes('1 item selecionado')")
            wait("q('review-data').textContent.includes('Liberado para nova tentativa')")
            assert load_queue(queue, reference_mode='direct').iloc[6]['Status'] == 'PENDENTE'
            assert session.allow_api is False and not session.active
            assert session.validate_campaign()['selection']['selected_ids'] == ['7']
            window.evaluate_js("q('review-section').scrollIntoView({block:'start'})")
            layout('review-released-feedback-small')
            window.evaluate_js("q('execution-selection').scrollIntoView({block:'start'})")
            layout('review-retry-selection-small')
            window.resize(1200,800)
            wait('innerWidth>1100')
            window.evaluate_js("q('review-section').scrollIntoView({block:'start'})")
            layout('review-released-feedback-wide')
            report['checks'].append('review-release-feedback-auto-select-only-authorized-no-generation')
            # Regression: an already authorized item can still say REVISAO in Excel.
            book = __import__('openpyxl').load_workbook(queue)
            book['Fila_Geracao']['F8'] = 'REVISAO'
            book.save(queue); book.close()
            click('#selection-preview')
            wait("q('selection-preview-data').textContent.includes('1 item selecionado')")
            assert session.validate_campaign()['selection']['selected_ids'] == ['7']
            window.evaluate_js("api=async function(path,body){if(path==='start'&&body?.real)throw Error('Real API forbidden');const result=await window.diagnosticOriginalApi(path,body);if(path==='state'){result.allow_api=true;result.api_key_status='session';}return result;};refresh()")
            wait('state.allow_api===true')
            click('#real-start')
            wait("q('generation-dialog').open && q('generation-items').textContent==='Até 1'")
            assert window.evaluate_js("q('generation-selected-ids').textContent.includes('7') && !q('generation-confirm').disabled")
            layout('authorized-review-modal-exact-id')
            click('#generation-cancel')
            window.evaluate_js('api=window.diagnosticOriginalApi;refresh()')
            wait('state.allow_api===false')
            assert session.allow_api is False and not session.active
            report['checks'].append('authorized-review-sheet-label-does-not-exclude-modal-id-no-api')
            # Large timeline fixture: never use the user's real 53-item campaign.
            from datetime import datetime, timezone
            stamp=datetime.now(timezone.utc).isoformat()
            events=[]
            for i in range(1,54):
                events.extend([dict(type='item_start',id=str(i),at=stamp),dict(type='item_end',id=str(i),at=stamp,
                    status='SIMULADO',elapsed_seconds=140,message='Registro fictício; nenhuma chamada à API.',
                    outputs=[str(root/'fictional-output'/f'produto-{i}.png')],api_attempts=0)])
            session.run_log=dict(version=2,run_id='timeline-fixture',started_at=stamp,finished_at=stamp,
                mode='simulation',status='Finalizado',planned_items=53,planned_images=53,elapsed_seconds=7466,
                mean_completed_seconds=140,events=events,counts=dict(planned=53,started=53,completed=53,errors=0,review=0,waiting=0))
            session.items=[dict(id=str(i),status='SIMULADO',message='Somente fixture') for i in range(1,54)]
            session.result=dict(success=53,processed=53,errors=0)
            session.total=53
            window.resize(1440,900)
            wait("innerWidth>1300 && q('run-history').querySelectorAll('tbody tr').length===106")
            wait("Math.abs(q('campaign-follow').getBoundingClientRect().bottom-q('campaign-setup').getBoundingClientRect().bottom)<2")
            window.evaluate_js("q('run-history').scrollIntoView({block:'center'});q('run-history').querySelector('.log-table-wrap').focus()")
            assert window.evaluate_js("(()=>{const e=q('run-history').querySelector('.log-table-wrap');return e.scrollHeight>e.clientHeight && e.clientHeight>=139;})()")
            window.evaluate_js("const timeline=q('run-history').querySelector('.log-table-wrap');timeline.scrollTop=timeline.scrollHeight")
            assert window.evaluate_js("q('run-history').querySelector('.log-table-wrap').scrollTop>0")
            assert window.evaluate_js("(()=>{const e=q('run-history').querySelector('.log-table-wrap');return Math.abs(e.querySelector('th').getBoundingClientRect().top-e.getBoundingClientRect().top)<4;})()")
            layout('timeline-53-wide-last-item')
            window.resize(520,500)
            wait('innerWidth<520')
            window.evaluate_js("q('run-history').querySelector('.log-table-wrap').scrollTop=0;q('run-history').querySelector('.log-table-wrap').scrollIntoView({block:'start'})")
            assert window.evaluate_js("(()=>{const e=q('run-history').querySelector('.log-table-wrap');return e.scrollHeight>e.clientHeight && e.clientHeight<=innerHeight*.6+2;})()")
            layout('timeline-53-small-scroll')
            window.resize(1440,900)
            wait('innerWidth>1300')
            window.evaluate_js("q('history-panel').open=false")
            wait("q('run-history').getBoundingClientRect().height===0")
            window.evaluate_js("q('history-panel').open=true;q('campaign-setup').style.paddingBottom='120px'")
            wait("Math.abs(q('campaign-follow').getBoundingClientRect().bottom-q('campaign-setup').getBoundingClientRect().bottom)<2")
            window.evaluate_js("q('campaign-setup').style.paddingBottom='';q('run-history').scrollIntoView({block:'center'})")
            layout('timeline-53-wide-restored')
            assert session.allow_api is False and not session.active
            report['checks'].append('timeline-53-internal-scroll-sticky-header-panel-alignment-small-and-resize')
            # Fluid workspace with the large timeline still present. Inspect
            # native client dimensions, not merely the requested outer size.
            report['responsive_sizes'] = []
            for requested in (520,768,1024,1366,1920,2560):
                window.resize(requested,900)
                time.sleep(.25)
                actual=window.evaluate_js('innerWidth')
                report['responsive_sizes'].append(dict(requested=requested,actual=actual,
                    limited_by_display=actual<requested-40))
                for theme in ('light','dark'):
                    window.evaluate_js("document.documentElement.dataset.theme="+json.dumps(theme)+";themeLabel();window.scrollTo(0,0)")
                    wait("(()=>{const l=q('campaign-setup').getBoundingClientRect(),r=q('campaign-follow').getBoundingClientRect();return innerWidth<=1100||Math.abs(l.bottom-r.bottom)<2})()")
                    bounds=window.evaluate_js("(()=>{const ids=['afirma-simples','main-content','campaign-grid','persistent-history'];return Object.fromEntries(ids.map(id=>{const r=q(id).getBoundingClientRect();return [id,{left:r.left,right:r.right,width:r.width}]}));})()")
                    root_bounds=bounds['afirma-simples']
                    if actual>=1800:
                        assert 24<=root_bounds['left']<=(40 if actual<=2000 else 48), bounds
                        assert root_bounds['width']>=actual*.94,bounds
                    for id in ('main-content','campaign-grid','persistent-history'):
                        assert abs(bounds[id]['left']-root_bounds['left'])<2,bounds
                        assert abs(bounds[id]['right']-root_bounds['right'])<2,bounds
                    layout(f'fluid-home-{requested}-{theme}')
                    window.evaluate_js("q('run-history').querySelector('.log-table-wrap').scrollIntoView({block:'start'})")
                    assert window.evaluate_js("(()=>{const e=q('run-history').querySelector('.log-table-wrap');return e.clientHeight>0 && e.scrollHeight>e.clientHeight})()")
                    layout(f'fluid-timeline-{requested}-{theme}')
                click('#settings-button')
                assert window.evaluate_js("q('settings').querySelector('.settings-panel').getBoundingClientRect().width<=1041")
                window.evaluate_js('window.scrollTo(0,0)')
                layout(f'fluid-settings-{requested}')
                click('#back')
            window.resize(1920,1080)
            time.sleep(.25)
            window.maximize()
            time.sleep(.25)
            window.evaluate_js("document.documentElement.dataset.theme='light';themeLabel();window.scrollTo(0,0)")
            layout('fluid-maximized-light')
            window.evaluate_js("document.documentElement.dataset.theme='dark';themeLabel()")
            layout('fluid-maximized-dark')
            window.restore();window.resize(1920,1080)
            time.sleep(.25)
            for factor in (1.25,1.5):
                on_ui_thread(window,lambda:setattr(window.native.browser.webview,'ZoomFactor',factor))
                time.sleep(.25)
                window.evaluate_js('window.scrollTo(0,0)')
                layout(f'fluid-wide-zoom-{int(factor*100)}')
                window.evaluate_js("q('run-history').querySelector('.log-table-wrap').scrollIntoView({block:'start'})")
                layout(f'fluid-timeline-zoom-{int(factor*100)}')
                click('#settings-button')
                window.evaluate_js('window.scrollTo(0,0)')
                layout(f'fluid-settings-zoom-{int(factor*100)}')
                click('#back')
            on_ui_thread(window,lambda:setattr(window.native.browser.webview,'ZoomFactor',1.0))
            window.restore();window.resize(1440,900)
            wait('innerWidth>1300')
            # Preserve the starting theme expected by subsequent flow checks.
            window.evaluate_js("document.documentElement.dataset.theme='light';themeLabel()")
            report['checks'].append('fluid-width-matrix-themes-aligned-edges-readable-settings-timeline-native-resize')
            # Home indicator: explicit fake probes only, shared state with Settings.
            session.set_selection('pending')
            window.evaluate_js("showSettings(false);window.scrollTo(0,0)")
            wait("q('home-api-label').textContent==='API desconectada'")
            layout('api-home-missing-light')
            session.set_api_key('fictional-home-indicator-only')
            wait("q('home-api-label').textContent==='API não verificada'")
            with patch('src.api_connection.check_api_connection') as probe:
                click('#home-api-status')
                assert window.evaluate_js("!q('settings').classList.contains('hidden') && document.querySelector('#settings-generation details').open")
                probe.assert_not_called()
                click('#back')
                window.evaluate_js("window.scrollTo(0,0);q('home-api-status').focus()")
                layout('api-home-unverified-keyboard-light')
                probe.assert_not_called()
            finish_check = threading.Event()
            def indicator_probe(key):
                finish_check.wait(10)
                return dict(status='connected', message='Autenticação fictícia confirmada.')
            with patch('src.api_connection.check_api_connection', side_effect=indicator_probe):
                click('#home-api-status');click('#api-key-check');click('#back')
                wait("q('home-api-label').textContent==='Verificando API…'")
                window.evaluate_js('window.scrollTo(0,0)')
                layout('api-home-checking-light')
                finish_check.set()
                wait("q('home-api-label').textContent==='API conectada'")
            layout('api-home-connected-light')
            for result, label in [(dict(status='unverified',category='limited',message='Limite fictício, verificação inconclusiva.'),'API não verificada'),
                                  (dict(status='disconnected',category='authentication',message='Autenticação fictícia recusada.'),'API desconectada')]:
                with patch('src.api_connection.check_api_connection', return_value=result):
                    click('#home-api-status');click('#api-key-check')
                    wait("state.api_connection.status==="+json.dumps(result['status']))
                    click('#back')
                    wait("q('home-api-label').textContent==="+json.dumps(label))
            click('#theme-toggle')
            window.evaluate_js('window.scrollTo(0,0)')
            layout('api-home-failed-dark')
            session.set_api_key('fictional-home-replacement-only')
            wait("q('home-api-label').textContent==='API não verificada'")
            layout('api-home-replaced-dark')
            session.clear_api_key('session')
            wait("q('home-api-label').textContent==='API desconectada'")
            report['checks'].append('home-api-states-explicit-only-synchronized-no-quota-as-disconnection')

            # Hold a fake item in flight so pause/stop requests can be inspected.
            entered, advance, resumed, finish_run = (threading.Event() for _ in range(4))
            def fake_run(*args, control, **kwargs):
                entered.set();advance.wait(30)
                control.checkpoint();resumed.set();finish_run.wait(30)
                return dict(success=0,errors=0,processed=0,stop_reason='Interrompido')
            with patch('main.main', side_effect=fake_run) as run:
                try:
                    window.evaluate_js("q('start').click();q('start').click()")
                    assert entered.wait(5)
                    wait("state.active && q('start').textContent==='Simulando campanha…'")
                    assert run.call_count==1
                    assert window.evaluate_js("[...document.querySelectorAll('.campaign-control')].every(e=>e.disabled) && !q('settings-button').disabled && !q('open-output').disabled")
                    window.evaluate_js("q('campaign-follow').scrollIntoView({block:'start'})")
                    layout('buttons-running-dark')
                    # Presentation of a real run, while the backend remains the
                    # same controlled fake worker and the API gate stays closed.
                    session.run_log['mode']='real'
                    window.evaluate_js("api=async function(path,body){if(path==='start'&&body?.real)throw Error('Real API forbidden');const result=await window.diagnosticOriginalApi(path,body);if(path==='state')result.allow_api=true;return result;};refresh()")
                    wait("q('real-start').disabled && q('real-start').textContent.includes('Gerando imagens')")
                    window.evaluate_js("q('real-start').scrollIntoView({block:'center'})")
                    layout('buttons-generating-fixture-dark')
                    session.run_log['mode']='simulation'
                    window.evaluate_js('api=window.diagnosticOriginalApi;refresh()')
                    click('#settings-button')
                    assert window.evaluate_js("q('save-settings').disabled && q('apply').disabled && q('max-jobs').disabled && !q('execution-lock-notice').classList.contains('hidden')")
                    window.evaluate_js('window.scrollTo(0,0)')
                    layout('settings-locked-dark')
                    click('#back')
                    (root/'images').mkdir(exist_ok=True)
                    with patch('src.studio_session.os.startfile',create=True) as opened:
                        click('#open-output')
                        wait("!polling")
                        deadline=time.monotonic()+3
                        while not opened.called and time.monotonic()<deadline:time.sleep(.05)
                        assert opened.call_count==1
                    click('#pause')
                    wait("state.pause_requested && !state.paused && q('pause').textContent.includes('Aguardando pausa')")
                    assert window.evaluate_js("q('pause').disabled && q('resume').classList.contains('hidden')")
                    window.evaluate_js("q('pause').scrollIntoView({block:'center'})")
                    layout('buttons-pause-requested-dark')
                    advance.set()
                    wait("state.paused && state.status==='Pausado'")
                    assert window.evaluate_js("!q('resume').disabled && q('pause').classList.contains('hidden')")
                    layout('buttons-paused-dark')
                    click('#resume')
                    assert resumed.wait(5)
                    wait('!state.paused && !state.pause_requested')
                    click('#stop')
                    wait("q('stop-dialog').open")
                    assert window.evaluate_js("document.activeElement.id==='stop-cancel'")
                    layout('buttons-stop-confirmation-dark')
                    click('#stop-cancel')
                    assert not session.control.stopped
                    click('#stop');click('#stop-confirm')
                    wait("state.stopping && q('stop').textContent.includes('Interrompendo')")
                    assert window.evaluate_js("q('stop').disabled && q('pause').disabled && q('start').disabled")
                    window.resize(520,600)
                    wait('innerWidth<520')
                    window.evaluate_js("q('stop').scrollIntoView({block:'center'})")
                    layout('buttons-stopping-small-dark')
                    click('#theme-toggle')
                    window.evaluate_js("q('stop').scrollIntoView({block:'center'})")
                    layout('buttons-stopping-small-light')
                finally:
                    advance.set();finish_run.set()
                    if session.control: session.control.command('stop')
                    if session.worker: session.worker.join(5)
            wait('!state.active')
            assert window.evaluate_js("!q('start').disabled && q('pause').disabled && q('stop').disabled")
            report['checks'].append('buttons-double-start-locks-consult-settings-folder-pause-effective-stop-confirmation')
            window.resize(1440,900)
            wait('innerWidth>1300')
            window.evaluate_js("api=async function(path,body){if(path==='start'&&body?.real)throw Error('Real API forbidden');const result=await window.diagnosticOriginalApi(path,body);if(path==='state')result.allow_api=true;return result;};refresh()")
            wait("!q('real-start').classList.contains('hidden')")
            window.evaluate_js("q('real-start').scrollIntoView({block:'center'})")
            layout('buttons-idle-light')
            click('#theme-toggle')
            window.evaluate_js("q('real-start').scrollIntoView({block:'center'})")
            layout('buttons-idle-dark')
            window.evaluate_js('api=window.diagnosticOriginalApi;refresh()')
            assert session.allow_api is False
            report["passed"] = True
        except Exception as exc:
            report["error"] = f"{type(exc).__name__}: {exc}"
            import traceback
            report['failure_line'] = traceback.extract_tb(exc.__traceback__)[-1].lineno
            report['failure_stack'] = [f'{frame.name}:{frame.lineno}' for frame in traceback.extract_tb(exc.__traceback__)]
            report['timeline_layout'] = window.evaluate_js("(()=>{const e=q('run-history')?.querySelector('.log-table-wrap');return {width:innerWidth,height:innerHeight,left:q('campaign-setup')?.getBoundingClientRect().height,right:q('campaign-follow')?.getBoundingClientRect().height,tableHeight:e?.clientHeight,tableScroll:e?.scrollHeight,scrollTop:e?.scrollTop,tableTop:e?.getBoundingClientRect().top,headerTop:e?.querySelector('th').getBoundingClientRect().top};})()")
            report['execution_ui'] = window.evaluate_js("({history:q('history-data')?.textContent,historyOptions:q('history-choice')?.innerHTML,runs:state.history?.runs?.length,selection:q('selection-preview-data')?.textContent,error:q('error')?.textContent})")
            report["ui_failure"] = window.evaluate_js("({dialogOpen:document.querySelector('#generation-dialog')?.open,error:document.querySelector('#generation-error')?.textContent,confirmDisabled:document.querySelector('#generation-confirm')?.disabled,keyStatus:state?.api_key_status,paidCalls:window.diagnosticPaidCalls})")
        finally:
            (root / "diagnostics.json").write_text(json.dumps(report, ensure_ascii=True, indent=2), "utf-8")
            window.destroy()

    return session, ready
