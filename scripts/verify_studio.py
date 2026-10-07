"""Create a reproducible fictitious campaign and serve the real Studio locally.

No .env, production workbook or network client is read. Run from project root.
"""
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PIL import Image
from openpyxl import Workbook
from src.studio_session import StudioSession
from studio import make_server


def fixture(root):
    root.mkdir(parents=True, exist_ok=True)
    queue = root/'campanha-ficticia.xlsx'
    if queue.exists():
        return queue  # Never overwrite an existing fixture or its possible history.
    for folder, color in [('produtos','purple'),('marca','orange')]:
        path = root / folder
        path.mkdir(exist_ok=True)
        Image.new('RGB',(160,120),color).save(path/'foto.png')
    book = Workbook()
    sheet = book.active
    sheet.title='Fila_Geracao'
    sheet.append(['ID','Prompt_Padrao','Prompt_Variacao','Arquivo_Referencia','Nome_Saida','Status','Tentativas','Observacao'])
    for i in range(1,13):
        sheet.append([f'DEMO-{i:02d}','Fotografia de produto em fundo claro.',f'Composição de demonstração {i}.',
                      'Produtos::foto.png; Marca::foto.png',f'demo_{i:02d}.png','PENDENTE',0,''])
    config=book.create_sheet('Configuracao')
    for row in [('Parametro','Valor'),('Pasta_Referencias',str(root/'produtos')),
                ('Pasta_Resultados',str(root/'imagens')),('Limite_Por_Execucao',12),('Dry_Run','SIM')]:
        config.append(row)
    folders=book.create_sheet('Pastas_Referencias')
    folders.append(['Alias','Caminho'])
    folders.append(['Produtos',str(root/'produtos')])
    folders.append(['Marca',str(root/'marca')])
    book.save(queue)
    book.close()
    return queue


if __name__=='__main__':
    root=ROOT/'output'/'studio-verification'
    queue=fixture(root)
    session=StudioSession(root/'preferences.json',env={})
    session.choose_queue(queue)
    session.add_photos([root/'produtos'/'foto.png'])
    session.add_photos([root/'marca'/'foto.png'])
    preferences=session.preferences.copy()
    preferences.update(max_jobs=12,max_images=12,safe_mode=False)
    session.set_preferences(preferences)
    server=make_server(session,port=int(sys.argv[1]) if len(sys.argv)>1 else 0)
    url=f'http://127.0.0.1:{server.server_port}/'
    (root/'verification.json').write_text(json.dumps({'url':url,'queue':str(queue),'api_enabled':False},indent=2),encoding='utf-8')
    print(url,flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        if session.active:
            session.command('stop')
    finally:
        server.server_close()
